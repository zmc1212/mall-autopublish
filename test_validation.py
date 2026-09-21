import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from PIL import Image

spec = importlib.util.spec_from_file_location('seller', Path(__file__).with_name('千牛自动上架.py'))
seller = importlib.util.module_from_spec(spec)
spec.loader.exec_module(seller)


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        Image.new('RGB', (2, 2)).save(self.root / 'main.png')
        self.data = dict(zip(seller.REQUIRED, ['ID1', '标题', '类目', 'main.png', 10, 0]))

    def errors(self, **changes):
        data = {**self.data, **changes}
        return seller.validate_row(list(data), list(data.values()), 2, self.root)[1]

    def workbook(self, rows=None):
        p = self.root / 'input.xlsx'
        w = Workbook()
        w.active.title = '商品清单'
        w.active.append(list(self.data) + ['处理状态', '上架时间'])
        for row in rows or [list(self.data.values())]:
            w.active.append(row)
        log = w.create_sheet('处理日志')
        log.append(['时间', 'Excel行号', '商品标识', '结果', '失败字段', '页面提示', '详情'])
        log.cell(201, 1).number_format = '@'
        w.save(p)
        w.close()
        return p

    def test_single_spec_relative_image(self):
        self.assertEqual(self.errors(), [])

    def test_numbers(self):
        for value in [float('nan'), float('inf'), -1, 0, True, 'x']:
            with self.subTest(price=value):
                self.assertTrue(self.errors(**{'价格*': value}))
        for value in [-1, 0.5, float('inf'), True]:
            with self.subTest(stock=value):
                self.assertTrue(self.errors(**{'库存*': value}))

    def test_json_and_skus(self):
        for field, value in [('商品属性(JSON)', '[]'), ('SKU明细(JSON)*', '{}'), ('规格1值(JSON)', '{')]:
            self.assertTrue(self.errors(**{field: value}))
        sku = {'规格': {'颜色': '红'}, '价格': 10, '库存': 1}
        fields = {'规格1名称': '颜色', '规格1值(JSON)': '["红"]', 'SKU明细(JSON)*': json.dumps([sku])}
        self.assertEqual(self.errors(**fields), [])
        fields['SKU明细(JSON)*'] = json.dumps([sku, sku])
        self.assertTrue(any('组合重复' in e for e in self.errors(**fields)))
        fields['SKU明细(JSON)*'] = ''
        self.assertTrue(self.errors(**fields))

    def test_exact_decimal_numbers(self):
        self.assertTrue(self.errors(**{'库存*': '1.00000000000000001'}))
        self.assertTrue(self.errors(**{'库存*': '-0.0000000000000000001'}))
        self.assertTrue(self.errors(**{'库存*': '9007199254740992.1'}))
        self.assertEqual(self.errors(**{'价格*': '0.0000000000000000001'}), [])

    def test_conflicting_sku_columns(self):
        self.assertTrue(any('新旧SKU列' in e for e in self.errors(**{
            'SKU明细(JSON)': '[]', 'SKU明细(JSON)*': '[]'})))

    def test_neutral_pen_required_attributes(self):
        data = {
            **self.data,
            '类目*': '文具用品/文化用品/商务用品>>笔类/书写工具>>中性笔',
            '商品标题*': '点石制笔中性笔考试专用',
            '品牌': '点石制笔',
            '发货时间': '48小时内发货',
            '发货地': '大陆及港澳台',
            '运费模板': '默认运费模板',
        }
        errors = seller.validate_row(list(data), list(data.values()), 2, self.root)[1]
        self.assertTrue(any('中性笔类目必须提供' in e for e in errors))
        attrs = {k: '值' for k in seller.NEUTRAL_PEN_REQUIRED_ATTRIBUTES}
        attrs['品牌'] = '点石制笔'
        data['商品属性(JSON)'] = json.dumps(attrs, ensure_ascii=False)
        self.assertEqual(seller.validate_row(list(data), list(data.values()), 2, self.root)[1], [])

    def test_brand_and_title_and_spec3(self):
        data = {**self.data, '商品标题*': '普通标题', '品牌': '点石制笔', '商品属性(JSON)': json.dumps({'品牌': '点石制笔'}, ensure_ascii=False)}
        errors = seller.validate_row(list(data), list(data.values()), 2, self.root)[1]
        self.assertTrue(any('必须包含品牌属性' in e for e in errors))
        sku = {'规格': {'颜色': '红', '粗细': '0.5', '支数': '1支'}, '价格': 10, '库存': 1}
        fields = {
            '商品标题*': '点石制笔中性笔',
            '规格1名称': '颜色', '规格1值(JSON)': '["红"]',
            '规格2名称': '粗细', '规格2值(JSON)': '["0.5"]',
            '规格3名称': '支数', '规格3值(JSON)': '["1支"]',
            'SKU明细(JSON)': json.dumps([sku], ensure_ascii=False),
        }
        self.assertEqual(self.errors(**fields), [])

    def test_example_and_title_limit(self):
        self.assertTrue(any('示例数据禁止提交' in e for e in self.errors(**{'商品标识*': '示例-001'})))
        self.assertTrue(any('超过30个汉字' in e for e in self.errors(**{'商品标题*': '测' * 31})))

    def test_portrait_ratio(self):
        Image.new('RGB', (3, 4)).save(self.root / 'p.png')
        Image.new('RGB', (4, 3)).save(self.root / 'badp.png')
        self.assertEqual(self.errors(**{'3:4主图路径': 'p.png'}), [])
        self.assertTrue(any('3:4' in e for e in self.errors(**{'3:4主图路径': 'badp.png'})))

    def test_images_and_warehouse(self):
        for name in ['missing.png', '.', 'bad.png']:
            (self.root / 'bad.png').write_text('not an image')
            self.assertTrue(self.errors(**{'主图路径*': name}))
        self.assertTrue(self.errors(**{'详情图路径': 'missing.png'}))
        self.assertTrue(self.errors(**{'上架时间': '立即上架'}))

    def test_blank_defaults_and_duplicate_ids(self):
        p = self.workbook([list(self.data.values()), [None]*6 + ['待处理', '放入仓库'], list(self.data.values())])
        results = seller.validate_workbook(p)
        self.assertEqual(len(results), 2)
        self.assertTrue(all(any('整批重复' in e for e in r['errors']) for r in results))

    def test_missing_columns(self):
        p = self.workbook()
        w = load_workbook(p)
        w.active.cell(1, 1, '错误列')
        w.save(p)
        w.close()
        with self.assertRaises(ValueError):
            seller.validate_workbook(p)

    def test_log_compaction_and_input_unchanged(self):
        p = self.workbook()
        before = p.read_bytes()
        out = seller.write_log(p, [['t', 2, 'ID1', '校验通过']])
        w = load_workbook(out)
        self.assertEqual(w['处理日志'].cell(2, 3).value, 'ID1')
        self.assertEqual(w['处理日志'].max_row, 2)
        w.close()
        out2 = seller.write_log(out, [['t2', 2, 'ID1', '校验通过']], self.root / 'second.xlsx')
        w = load_workbook(out2)
        self.assertEqual(w['处理日志'].max_row, 3)
        w.close()
        self.assertEqual(before, p.read_bytes())
        with self.assertRaises(ValueError):
            seller.write_log(p, [], p)

    def test_write_log_strips_illegal_excel_characters(self):
        p = self.workbook()
        nasty = "TimeoutError:\x1b[2m waiting\x1b[22m \x00 \x08 overlay"
        out = seller.write_log(p, [['t', 2, 'ID1', '失败', '', '', nasty]])
        w = load_workbook(out)
        detail = w['处理日志'].cell(2, 7).value
        w.close()
        self.assertNotIn('\x1b', detail)
        self.assertNotIn('\x00', detail)
        self.assertIn('TimeoutError', detail)
        self.assertIn('overlay', detail)

    def test_write_log_includes_taobao_item_id(self):
        p = self.workbook()
        out = seller.write_log(
            p,
            [['t', 2, 'ID1', '结果待核实', '', '提交成功', '', '1086638256748']],
        )
        w = load_workbook(out)
        log = w['处理日志']
        self.assertEqual(log.cell(1, 8).value, '淘宝商品ID')
        self.assertEqual(log.cell(2, 8).value, '1086638256748')
        w.close()

    def test_failed_excel_retains_report(self):
        p = self.workbook()
        report = self.root / 'report.json'
        with patch('sys.argv', ['seller', str(p), '--log', str(report)]), patch.object(seller, 'write_log', side_effect=PermissionError('locked')):
            self.assertEqual(seller.main(), 2)
        self.assertEqual(json.loads(report.read_text(encoding='utf-8'))['results'][0]['execution'], '未执行')

    def test_submit_fills_without_confirm(self):
        p = self.workbook()
        executed = [{'row': 2, 'execution': '已填写未提交', 'notice': '已选择放入仓库', 'errors': []}]
        with patch('sys.argv', ['seller', str(p), '--submit']), patch.object(seller, 'run_web_batch', return_value=executed) as mocked:
            self.assertEqual(seller.main(), 0)
            mocked.assert_called_once()
            self.assertFalse(mocked.call_args.kwargs['confirm_submit'])

    def test_example_row_not_sent_to_browser(self):
        p = self.workbook([['示例-001', '示例商品标题', '类目', 'main.png', 10, 0]])
        with patch('sys.argv', ['seller', str(p), '--submit', '--confirm-submit']), patch.object(seller, 'run_web_batch') as mocked:
            self.assertEqual(seller.main(), 1)
            mocked.assert_not_called()


def naruto_pack():
    testdata = Path(__file__).resolve().parent / 'testdata'
    for path in testdata.iterdir():
        if path.is_dir() and '火影' in path.name:
            return path
    raise FileNotFoundError('找不到火影图片包')


class ImagePackAndTemplateTests(unittest.TestCase):
    def test_scan_naruto_pack(self):
        pack = seller.scan_image_pack(naruto_pack())
        self.assertEqual(len(pack['main_1_1']), 5)
        self.assertEqual(len(pack['details']), 7)
        self.assertEqual(len(pack['skus']), 16)
        self.assertEqual(pack['main_3_4'], [])
        self.assertEqual([p.name for p in pack['main_1_1']], [f'宝贝主图0{i}.jpg' for i in range(1, 6)])
        self.assertEqual(pack['details'][-1].name, '详情07.png')
        self.assertEqual(pack['skus'][0]['slot'], '颜色01')
        self.assertEqual(pack['skus'][0]['name'], '忍道版1弹【特别款】1支')
        self.assertEqual(pack['skus'][-1]['slot'], '颜色16')
        self.assertEqual(pack['skus'][-1]['name'], '忍道版1弹【经典款】李洛克')

    def test_merge_templates_missing_brand(self):
        templates = seller.load_templates({'attributes': '中性笔', 'logistics': '48小时', 'sales': '仓库多规格'})
        self.assertNotIn('品牌', templates['attributes']['attributes'])
        self.assertNotIn('型号', templates['attributes']['attributes'])
        self.assertEqual(templates['logistics']['ship_time'], '48小时内发货')
        self.assertEqual(templates['logistics']['freight'], '文具用品 包邮')
        product = seller.resolve_product({
            '商品标题*': '火影忍者中性笔盲盒',
            '型号': '忍道版第1弹',
            '价格*': 9.9,
            '库存*': 10,
        }, [], None, templates)
        self.assertEqual(product['category'], '文具用品/文化用品/商务用品>>笔类/书写工具>>中性笔')
        self.assertFalse(product.get('brand'))
        errors = seller.validate_product(product)
        self.assertTrue(any('品牌' in e for e in errors))

    def test_template_row_does_not_need_attribute_json(self):
        pack = naruto_pack()
        data = {
            '商品标识*': '火影-001',
            '商品标题*': '卡游火影忍者中性笔盲盒忍道版',
            '品牌*': '卡游',
            '型号': '忍道版第1弹',
            '图片包路径*': str(pack),
            '商品属性模板*': '中性笔',
            '物流模板*': '48小时',
            '销售模板*': '仓库多规格',
            '价格*': 9.9,
            '库存*': 20,
        }
        errors = seller.validate_row(list(data), list(data.values()), 2, pack.parent)[1]
        self.assertEqual(errors, [])
        self.assertFalse(any('商品属性(JSON)' in e for e in errors))
        product = seller.row_to_product(data, pack.parent)
        self.assertEqual(product['brand'], '卡游')
        self.assertEqual(product['attributes']['笔头类型'], '子弹头')
        self.assertEqual(product['attributes']['品牌'], '卡游')
        self.assertEqual(product['attributes']['型号'], '忍道版第1弹')
        self.assertEqual(product['ship_time'], '48小时内发货')
        self.assertEqual(product['listing_time'], '放入仓库')
        self.assertEqual(len(product['main_images']), 5)
        self.assertEqual(len(product['detail_images']), 7)
        self.assertEqual(len(product['skus']), 16)
        self.assertEqual(product['skus'][0]['规格']['商品规格'], '忍道版1弹【特别款】1支')
        self.assertEqual(product['skus'][0]['价格'], 9.9)
        self.assertEqual(product['thickness'], '0.05mm')
        self.assertEqual(product['sku_category'], '单品')
        self.assertTrue(str(product.get('pack_dir') or '').endswith(pack.name) or pack.name in str(product.get('pack_dir') or ''))

    def test_create_template_dropdowns_and_naruto_row(self):
        root = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__('shutil').rmtree(root, ignore_errors=True))
        path = seller.create_template(root / 'tpl.xlsx')
        wb = load_workbook(path)
        self.assertIn('商品清单', wb.sheetnames)
        self.assertIn('SKU规格', wb.sheetnames)
        self.assertIn('_模板选项', wb.sheetnames)
        self.assertIn('字段说明', wb.sheetnames)
        self.assertEqual(wb['_模板选项'].sheet_state, 'hidden')
        self.assertEqual(wb['_模板选项']['A1'].value, '中性笔')
        self.assertEqual([wb['_模板选项']['B1'].value, wb['_模板选项']['B2'].value], ['48小时', '24小时'])
        self.assertEqual(wb['_模板选项']['C1'].value, '仓库多规格')
        headers = [cell.value for cell in wb['商品清单'][1]]
        self.assertEqual(headers, seller.TEMPLATE_HEADERS)
        row = dict(zip(headers, [cell.value for cell in wb['商品清单'][2]]))
        self.assertEqual(row['商品标识*'], '火影-001')
        self.assertIn('卡游', row['商品标题*'])
        self.assertLessEqual(seller.title_width(row['商品标题*']), 60)
        self.assertEqual(row['商品属性模板*'], '中性笔')
        self.assertEqual(row['物流模板*'], '48小时')
        self.assertEqual(row['销售模板*'], '仓库多规格')
        self.assertIn('火影', str(row['图片包路径*']))
        sku_headers = [cell.value for cell in wb['SKU规格'][1]]
        self.assertEqual(sku_headers, seller.SKU_HEADERS)
        self.assertEqual(wb['SKU规格'].max_row, 17)
        self.assertEqual(wb['SKU规格']['B2'].value, '颜色01')
        self.assertEqual(wb['SKU规格']['B17'].value, '颜色16')
        formulas = [rule.formula1 for rule in wb['商品清单'].data_validations.dataValidation]
        self.assertEqual(len(formulas), 3)
        wb.close()
        results = seller.validate_workbook(path)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['errors'], [])
        self.assertEqual(len(results[0]['product']['skus']), 16)

    def test_legacy_json_workbook_still_validates(self):
        root = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__('shutil').rmtree(root, ignore_errors=True))
        Image.new('RGB', (8, 8)).save(root / 'main.png')
        attrs = {k: '值' for k in seller.NEUTRAL_PEN_REQUIRED_ATTRIBUTES}
        attrs['品牌'] = '点石制笔'
        sku = {'规格': {'颜色': '红'}, '价格': 10, '库存': 1}
        headers = [
            '商品标识*', '商品标题*', '类目*', '主图路径*', '价格*', '库存*',
            '品牌', '发货时间', '发货地', '运费模板', '商品属性(JSON)',
            '规格1名称', '规格1值(JSON)', 'SKU明细(JSON)',
        ]
        values = [
            'OLD-001', '点石制笔中性笔考试专用',
            '文具用品/文化用品/商务用品>>笔类/书写工具>>中性笔',
            'main.png', 10, 2, '点石制笔', '48小时内发货', '大陆及港澳台', '默认运费模板',
            json.dumps(attrs, ensure_ascii=False), '颜色', '["红"]', json.dumps([sku], ensure_ascii=False),
        ]
        path = root / 'legacy.xlsx'
        wb = Workbook()
        wb.active.title = '商品清单'
        wb.active.append(headers)
        wb.active.append(values)
        wb.save(path)
        wb.close()
        results = seller.validate_workbook(path)
        self.assertEqual(results[0]['errors'], [])
        self.assertEqual(results[0]['product']['brand'], '点石制笔')
        self.assertEqual(results[0]['product']['skus'][0]['规格']['颜色'], '红')
        self.assertFalse(results[0]['product'].get('thickness'))

    def test_sku_sheet_overrides_filename(self):
        pack = seller.scan_image_pack(naruto_pack())
        templates = seller.load_templates({'attributes': '中性笔', 'logistics': '48小时', 'sales': '仓库多规格'})
        row = {
            '商品标识*': '火影-001',
            '商品标题*': '卡游火影忍者中性笔盲盒忍道版',
            '品牌*': '卡游',
            '型号': '忍道版第1弹',
            '价格*': 9.9,
            '库存*': 20,
        }
        sku_rows = [{'商品标识': '火影-001', '色号': '颜色01', '规格名称': '改名规格', '价格': 8.8, '库存': 3}]
        product = seller.resolve_product(row, sku_rows, pack, templates)
        self.assertEqual(product['skus'][0]['name'], '改名规格')
        self.assertEqual(product['skus'][0]['价格'], 8.8)
        self.assertEqual(product['skus'][0]['库存'], 3)
        self.assertEqual(product['skus'][1]['name'], pack['skus'][1]['name'])
        self.assertEqual(product['skus'][1]['价格'], 9.9)

    def test_unknown_template_name_is_error(self):
        data = {
            '商品标识*': '火影-001',
            '商品标题*': '卡游火影忍者中性笔盲盒忍道版',
            '品牌*': '卡游',
            '图片包路径*': str(naruto_pack()),
            '商品属性模板*': '不存在的模板',
            '物流模板*': '48小时',
            '销售模板*': '仓库多规格',
            '价格*': 9.9,
            '库存*': 20,
        }
        errors = seller.validate_row(list(data), list(data.values()), 2, naruto_pack().parent)[1]
        self.assertTrue(any('未知' in e and '商品属性' in e for e in errors))

    def test_portrait_filename_classified(self):
        root = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__('shutil').rmtree(root, ignore_errors=True))
        Image.new('RGB', (3, 4)).save(root / '主图3比4-01.jpg')
        Image.new('RGB', (8, 8)).save(root / '宝贝主图01.jpg')
        pack = seller.scan_image_pack(root)
        self.assertEqual([p.name for p in pack['main_1_1']], ['宝贝主图01.jpg'])
        self.assertEqual([p.name for p in pack['main_3_4']], ['主图3比4-01.jpg'])


if __name__ == '__main__':
    unittest.main()
