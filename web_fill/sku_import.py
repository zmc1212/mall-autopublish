"""Build a per-product SKU import file from the seller's category template."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4
import warnings

from openpyxl import load_workbook


CATEGORY_ID = "50012720"
HEADERS = ("颜色分类", "书写粗细", "支数", "商品规格", "价格", "数量", "商家编码", "商品条形码")


def supports_product(payload: dict) -> bool:
    return (
        bool(payload.get("skus"))
        and str(payload.get("category_id") or "") == CATEGORY_ID
        and str(payload.get("spec_name") or "商品规格") == "商品规格"
    )


def build_import_file(payload: dict, template: Path, directory: Path) -> Path:
    """Keep the original OOXML structure, including its hidden worksheet and .xls suffix."""
    if not supports_product(payload):
        raise ValueError("当前商品类目或规格列与 SKU 导入模板不匹配")
    if not template.is_file():
        raise FileNotFoundError(f"缺少千牛 SKU 导入模板: {template}")
    # The file supplied by Qianniu is OOXML despite its .xls name. Loading bytes
    # avoids openpyxl's extension guard while retaining the seller's file suffix.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(BytesIO(template.read_bytes()))
    sheet = workbook.worksheets[0]
    actual = tuple(sheet.cell(1, index).value for index in range(1, len(HEADERS) + 1))
    if actual != HEADERS:
        raise ValueError(f"千牛 SKU 模板表头已变化: {actual}")
    if sheet.max_row != 1:
        raise ValueError("千牛 SKU 模板包含数据行，已停止生成以免覆盖")
    for index, sku in enumerate(payload["skus"], 1):
        name = str(sku.get("name") or "").strip()
        if not name:
            raise ValueError(f"第 {index} 个 SKU 缺少商品规格")
        price = sku.get("price")
        stock = sku.get("stock")
        if price in (None, "") or stock in (None, ""):
            raise ValueError(f"SKU {name} 缺少价格或数量")
        # 书写粗细是非必填列：留空，不进模板也不在页面上点选，
        # 避免触发滑块认证与多维度识别的"在当前规格后添加"卡顿。
        sheet.append([
            "", "", "", name,
            price, stock, str(sku.get("merchant_code") or "").strip(),
            str(sku.get("barcode") or "").strip(),
        ])
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / f"SKU导入_{uuid4().hex}.xls"
    workbook.save(destination)
    return destination
