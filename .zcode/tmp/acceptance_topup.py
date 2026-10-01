# -*- coding: utf-8 -*-
"""真实页面验收：在保留现场(17行)上执行 _topup_missing_sku_rows 的生产路径。
skus.js 逐值补齐缺失规格 -> sku_import.js verify 复核 -> 报告行数。"""
import io
import json
import sys
import warnings

sys.path.insert(0, ".")
sys.path.insert(0, "desktop")
warnings.filterwarnings("ignore")

from desktop import paths

paths.configure_environ()

import 千牛网页执行 as web
from web_fill import pipeline

TEMPLATE_DIR = paths.load_settings().results_dir
TEMPLATE = None
import glob
import os

cands = sorted(glob.glob(os.path.join(TEMPLATE_DIR, "sku-import", "SKU导入_*.xls")), key=os.path.getmtime)
TEMPLATE = cands[-1]
print("template:", os.path.basename(TEMPLATE))

# 1. 从模板构建 payload（与 sku_import.build_import_file 的列序一致：4=名称 5=价格 6=数量 7=商家编码 8=条码）
from io import BytesIO

from openpyxl import load_workbook

wb = load_workbook(BytesIO(open(TEMPLATE, "rb").read()))
ws = wb.worksheets[0]
skus = []
for r in range(2, ws.max_row + 1):
    name = str(ws.cell(r, 4).value or "").strip()
    if not name:
        continue
    skus.append({
        "name": name,
        "price": ws.cell(r, 5).value,
        "stock": ws.cell(r, 6).value,
        "merchant_code": str(ws.cell(r, 7).value or "").strip(),
        "barcode": str(ws.cell(r, 8).value or "").strip(),
    })
print("skus:", len(skus))
payload = {"skus": skus, "sku_category": "单品", "thickness": ""}

# 2. 会话
session = web.CliSession()
session.attach()
print("attached, href:", session.href()[:80])

steps = []


def row_count():
    return session.eval_json(
        "() => ({n: [...document.querySelectorAll('table tr')].filter(tr => {"
        " const t = tr.innerText || '';"
        " return t.includes('元') && t.includes('件') && !t.includes('SKU分类');}).length})"
    ).get("n")


print("rows before:", row_count())

# 3. 生产补值路径：skus.js（skip_thickness）
try:
    result = pipeline.run_gate(session, "skus.js", {**payload, "skip_thickness": True},
                               steps, "sku_seed", timeout=600)
    print("skus.js done")
    print(json.dumps(result, ensure_ascii=False)[:1500])
except Exception as exc:
    print("skus.js FAILED:", str(exc)[:600])
    sys.exit(1)

# 4. 复核：sku_import.js verify
try:
    verify = pipeline.run_script(session, "sku_import.js",
                                 {"file": TEMPLATE, "skus": payload["skus"], "phase": "verify"},
                                 timeout=120)
    print("verify:", json.dumps(verify, ensure_ascii=False)[:600])
except Exception as exc:
    print("verify FAILED:", str(exc)[:400])

print("rows after:", row_count())
