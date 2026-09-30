import importlib
import importlib.util
import json
import os
import sys
import tempfile
import threading
import time
import unittest
import unittest.mock
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path


class DesktopModuleTests(unittest.TestCase):
    def test_concurrent_load_waits_for_module_initialization(self):
        from desktop.modules import load_named

        name = "_qianniu_slow_load_test"
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "slow_module.py"
            source.write_text("import time\ntime.sleep(0.2)\nready = True\n", encoding="utf-8")
            start = threading.Barrier(2)

            def load():
                start.wait()
                return load_named(name, source.name)

            try:
                with unittest.mock.patch("desktop.modules.find_resource", return_value=source):
                    with ThreadPoolExecutor(max_workers=2) as pool:
                        modules = list(pool.map(lambda _: load(), range(2)))
                self.assertIs(modules[0], modules[1])
                self.assertTrue(all(module.ready for module in modules))
            finally:
                sys.modules.pop(name, None)


class DesktopPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.appdata = Path(self.temp.name)
        os.environ["QIANNIU_APPDATA"] = str(self.appdata)
        os.environ.pop("QIANNIU_PROFILE", None)
        os.environ.pop("QIANNIU_OUTPUT_DIR", None)
        os.environ.pop("QIANNIU_CHROME", None)

    def test_profile_and_output_use_appdata(self):
        from desktop import paths
        importlib.reload(paths)
        settings = paths.configure_environ()
        self.assertTrue(str(Path(settings.chrome_profile)).startswith(str(self.appdata)))
        self.assertTrue(str(paths.playwright_output_dir()).startswith(str(self.appdata)))
        self.assertTrue(Path(settings.results_dir).is_dir())
        self.assertTrue(paths.cli_js_path().is_file())
        self.assertTrue(paths.templates_dir().is_dir())

    def test_web_module_reload_paths(self):
        from desktop.paths import configure_environ
        configure_environ()
        spec = importlib.util.spec_from_file_location(
            "web_cfg", Path(__file__).with_name("千牛网页执行.py")
        )
        web = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(web)
        web.reload_paths()
        self.assertTrue(str(web.PROFILE).startswith(str(self.appdata)))
        self.assertNotEqual(str(web.PROFILE), r"G:\workspace\taobao-playwright-profile")
        self.assertTrue(web.CHROME.is_file() or "chrome.exe" in str(web.CHROME).lower())
        args = web.chrome_launch_args()
        self.assertTrue(any(item.startswith("--user-data-dir=") for item in args))
        self.assertTrue(any(item.startswith("--remote-debugging-port=") for item in args))
        self.assertNotIn("https://myseller.taobao.com/", args)
        self.assertNotIn("--restore-last-session", args)
        prefs = Path(web.PROFILE) / "Default" / "Preferences"
        self.assertTrue(prefs.is_file())
        data = __import__("json").loads(prefs.read_text(encoding="utf-8"))
        self.assertEqual(data["session"]["restore_on_startup"], 1)

        cookies = Path(web.PROFILE) / "Default" / "Network"
        cookies.mkdir(parents=True, exist_ok=True)
        (cookies / "Cookies").write_bytes(b"cookie")
        restored = web.chrome_launch_args()
        self.assertIn("--restore-last-session", restored)
        self.assertNotIn("https://myseller.taobao.com/", restored)
        self.assertIn("https://myseller.taobao.com/", web.chrome_launch_args(url=web.SELLER_HOME))

    def test_invalid_custom_browser_falls_back_to_bundled(self):
        from desktop import paths

        bundled = self.appdata / "browser" / "chromium" / "chrome.exe"
        bundled.parent.mkdir(parents=True)
        bundled.write_bytes(b"mz")
        settings = paths.Settings(chrome_path=str(self.appdata / "missing-chrome.exe")).normalized()
        self.assertEqual(settings.chrome_path, "")
        with unittest.mock.patch("desktop.paths.bundled_browser_path", return_value=bundled), \
             unittest.mock.patch("desktop.paths.detect_system_chrome", return_value=None), \
             unittest.mock.patch.dict(os.environ, {"QIANNIU_CHROME": "", "CHROME_PATH": ""}):
            self.assertEqual(paths.browser_exe_path(settings.chrome_path), bundled)
            self.assertEqual(paths.browser_source(settings.chrome_path), "bundled")

    def test_confirm_submit_no_longer_a_setting(self):
        # 入库固定提交到仓库，不再作为可配置项
        from desktop import paths

        self.assertFalse(hasattr(paths.Settings(), "confirm_submit"))
        self.assertFalse("confirm_submit" in paths.load_settings().__dict__)
        self.assertFalse(paths.Settings().sku_template_import)
        self.assertFalse(paths.load_settings().sku_template_import)

    def test_sku_template_import_setting_persists(self):
        from desktop import paths

        paths.save_settings(paths.Settings(sku_template_import=True))
        self.assertTrue(paths.load_settings().sku_template_import)

    def test_spec_upload_batch_size_setting_persists_and_clamps(self):
        from desktop import paths

        # 默认每批 2 行；0 表示全部一次上传。
        self.assertEqual(paths.Settings().spec_upload_batch_size, 2)
        paths.save_settings(paths.Settings(spec_upload_batch_size=5))
        self.assertEqual(paths.load_settings().spec_upload_batch_size, 5)
        paths.save_settings(paths.Settings(spec_upload_batch_size=0))
        self.assertEqual(paths.load_settings().spec_upload_batch_size, 0)
        self.assertEqual(paths.Settings(spec_upload_batch_size=200).normalized().spec_upload_batch_size, 99)
        self.assertEqual(paths.Settings(spec_upload_batch_size=-3).normalized().spec_upload_batch_size, 0)

    def test_frozen_build_preserves_debug_browser_setting(self):
        from desktop import paths

        with unittest.mock.patch("desktop.paths.is_frozen", return_value=True):
            settings = paths.Settings(debug_browser=True).normalized()
        self.assertTrue(settings.debug_browser)


class DesktopApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        os.environ["QIANNIU_APPDATA"] = self.temp.name
        from desktop.paths import configure_environ
        configure_environ()

    def test_import_and_validate_sample_workbook(self):
        from fastapi.testclient import TestClient
        from desktop.server import app
        from desktop.jobs import MANAGER

        MANAGER._restored = True
        MANAGER.restored = False
        MANAGER.rows = []
        MANAGER.products = []
        MANAGER.workbook_path = ""
        excel = Path(__file__).resolve().parent / "testdata" / "中性笔测试入库.xlsx"
        self.assertTrue(excel.is_file(), excel)
        client = TestClient(app)
        response = client.post("/api/workbook/import", json={"path": str(excel)})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertGreaterEqual(data["count"], 1)
        self.assertGreaterEqual(data["valid"], 1)
        self.assertEqual(data["rows"][0]["validation"], "通过")
        health = client.get("/api/health")
        self.assertEqual(health.json()["ok"], True)
        status = client.get("/api/status")
        self.assertIn("chrome", status.json())
        self.assertIn("chrome_found", status.json()["chrome"])

    def test_settings_spec_upload_batch_size_roundtrip(self):
        from fastapi.testclient import TestClient
        from desktop.server import app

        client = TestClient(app)
        response = client.put("/api/settings", json={"spec_upload_batch_size": 3})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["spec_upload_batch_size"], 3)
        status = client.get("/api/status").json()
        self.assertEqual(status["settings"]["spec_upload_batch_size"], 3)
        # 越界值被 pydantic 拒绝（ge=0, le=99）
        self.assertEqual(client.put("/api/settings", json={"spec_upload_batch_size": 150}).status_code, 422)
        self.assertEqual(client.put("/api/settings", json={"spec_upload_batch_size": -1}).status_code, 422)

    def test_pending_products_skip_done_unless_force_new(self):
        from desktop.jobs import JobManager
        manager = JobManager()
        manager.workbook_path = "x.xlsx"
        manager.products = [
            {"meta": {"validation": "通过", "execution": "已填写未提交", "row": 2, "errors": []}, "product": {"row": 2, "product_id": "A"}},
            {"meta": {"validation": "通过", "execution": "失败", "row": 3, "errors": []}, "product": {"row": 3, "product_id": "B"}},
            {"meta": {"validation": "失败", "execution": "未执行", "row": 4, "errors": []}, "product": {"row": 4, "product_id": "C"}},
        ]
        manager.rows = [
            {"validation": "通过", "execution": "已填写未提交"},
            {"validation": "通过", "execution": "失败"},
            {"validation": "失败", "execution": "未执行"},
        ]
        self.assertEqual([item["product_id"] for item in manager._pending_products(False)], ["B"])
        self.assertEqual([item["product_id"] for item in manager._pending_products(True)], ["A", "B"])
        self.assertEqual([item["product_id"] for item in manager._pending_products(retry_failed=True)], ["B"])
        snap = manager.snapshot()
        self.assertTrue(snap["can_resume"])
        self.assertEqual(snap["pending"], 1)
        self.assertEqual(snap["retryable"], 1)

    def test_existing_item_id_is_never_filtered_as_legacy_done(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager.workbook_path = "x.xlsx"
        item = {
            "row": 2,
            "product_id": "A",
            "validation": "通过",
            # 旧记录可能仍保留提交阶段文案，但 ID 已证明商品已经入库。
            "execution": "结果待核实",
            "taobao_item_id": "1088898192279",
            "errors": [],
        }
        manager.products = [{"meta": dict(item), "product": {"row": 2, "product_id": "A"}}]
        manager.rows = [dict(item)]

        self.assertFalse(manager._row_flow_done(item))
        self.assertTrue(manager._row_flow_resumable(item))
        self.assertEqual([p["product_id"] for p in manager._pending_products(False)], ["A"])
        self.assertEqual([p["product_id"] for p in manager._pending_products(retry_failed=True)], ["A"])

    def test_pending_products_uses_execution_row_item_id_when_meta_lags(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager.workbook_path = "x.xlsx"
        manager.products = [{
            "meta": {
                "row": 2,
                "product_id": "A",
                "validation": "通过",
                "execution": "暂停",
            },
            "product": {"row": 2, "product_id": "A"},
        }]
        manager.rows = [{
            "row": 2,
            "product_id": "A",
            "validation": "通过",
            "execution": "暂停",
            "taobao_item_id": "1088898192279",
            "flow_stage": "material_verifying",
        }]

        selected = manager._pending_products(False)
        self.assertEqual(selected[0]["taobao_item_id"], "1088898192279")
        self.assertEqual(selected[0]["flow_stage"], "material_verifying")

    def test_restores_workbook_and_execution_after_restart(self):
        from desktop.jobs import JobManager
        excel = Path(__file__).resolve().parent / "testdata" / "中性笔测试入库.xlsx"
        first = JobManager()
        first._restored = True
        snap = first.import_workbook(str(excel))
        self.assertGreaterEqual(snap["count"], 1)
        first.products[0]["meta"]["execution"] = "失败"
        first.products[0]["meta"]["notice"] = "规格图中断"
        first.rows[0]["execution"] = "失败"
        first.rows[0]["notice"] = "规格图中断"
        first.status = "error"
        first._save_session()

        second = JobManager()
        restored = second.snapshot()
        self.assertTrue(restored["restored"])
        self.assertTrue(restored["workbook_path"].endswith(excel.name))
        self.assertGreaterEqual(restored["count"], 1)
        self.assertEqual(restored["rows"][0]["execution"], "失败")
        self.assertEqual(restored["rows"][0]["notice"], "规格图中断")
        self.assertGreaterEqual(restored["pending"], 1)

    def test_restores_from_latest_result_json(self):
        from desktop.jobs import JobManager
        from desktop import paths
        excel = Path(__file__).resolve().parent / "testdata" / "中性笔测试入库.xlsx"
        from openpyxl import load_workbook
        wb = load_workbook(excel, read_only=True, data_only=True)
        try:
            sheet = wb["商品清单"] if "商品清单" in wb.sheetnames else wb.worksheets[0]
            row_iter = sheet.iter_rows(values_only=True)
            headers = [str(cell or "").strip() for cell in next(row_iter)]
            product_id = str(dict(zip(headers, next(row_iter))).get("商品标识*") or "").strip()
        finally:
            wb.close()
        results_dir = Path(os.environ["QIANNIU_APPDATA"]) / "results"
        results_dir.mkdir(parents=True, exist_ok=True)
        journal = results_dir / "中性笔测试入库.结果-20260101-010101-1.json"
        journal.write_text(
            __import__("json").dumps({
                "source": str(excel.resolve()),
                "results": [{
                    "row": 2,
                    "product_id": product_id,
                    "execution": "失败",
                    "notice": "中断",
                    "validation": "通过",
                    "errors": [],
                }],
            }, ensure_ascii=False),
            encoding="utf-8",
        )
        os.environ["QIANNIU_RESULTS_DIR"] = str(results_dir)
        paths.configure_environ()
        manager = JobManager()
        snap = manager.snapshot()
        self.assertTrue(snap["restored"])
        self.assertTrue(snap["workbook_path"].endswith(excel.name))
        self.assertEqual(snap["rows"][0]["execution"], "失败")

    def test_restore_desktop_manager_fills_frozen_manager(self):
        import job_session
        from desktop.jobs import MANAGER
        from desktop import paths
        excel = Path(__file__).resolve().parent / "testdata" / "卡游火影忍者中性笔清单.xlsx"
        self.assertTrue(excel.is_file(), excel)
        results_dir = Path(os.environ["QIANNIU_APPDATA"]) / "results"
        results_dir.mkdir(parents=True, exist_ok=True)
        journal = results_dir / "卡游火影忍者中性笔清单.结果-20990101-010101-1.json"
        journal.write_text(
            __import__("json").dumps({
                "source": str(excel.resolve()),
                "results": [{
                    "row": 2,
                    "product_id": "火影-001",
                    "execution": "暂停",
                    "notice": "详情图中断",
                    "validation": "通过",
                    "errors": [],
                }],
            }, ensure_ascii=False),
            encoding="utf-8",
        )
        os.environ["QIANNIU_RESULTS_DIR"] = str(results_dir)
        paths.configure_environ()
        MANAGER.workbook_path = ""
        MANAGER.products = []
        MANAGER.rows = []
        MANAGER.status = "idle"
        job_session._RESTORE_TRIED = False
        self.assertTrue(job_session.restore_desktop_manager())
        self.assertTrue(MANAGER.workbook_path.endswith(excel.name))
        self.assertGreaterEqual(len(MANAGER.rows), 1)
        self.assertEqual(MANAGER.rows[0]["execution"], "暂停")
        self.assertTrue(MANAGER.products)

    def test_serialize_row_builds_item_buttons_from_notice(self):
        from desktop.jobs import serialize_row
        row = serialize_row({
            "row": 2,
            "product_id": "火影-001",
            "execution": "结果待核实",
            "notice": "商品提交成功 商品ID: 1086638256748, 您可以在商品列表中根据商品 ID搜索找到该商品。",
            "product": {"title": "火影"},
        })
        self.assertEqual(row["taobao_item_id"], "1086638256748")
        self.assertIn("item.htm?id=1086638256748", row["view_url"])
        self.assertIn("itemId=1086638256748", row["edit_url"])

    def test_open_item_opens_view_url(self):
        from desktop.jobs import JobManager
        manager = JobManager()
        manager._restored = True
        manager.rows = [{
            "row": 2,
            "product_id": "火影-001",
            "notice": "商品提交成功 商品ID: 1086638256748",
        }]
        opened = []

        class FakeWeb:
            def open_item_url(self, url):
                opened.append(url)
                return {"ok": True, "url": url, "via": "default-browser"}

        with unittest.mock.patch("desktop.jobs.load_web", return_value=FakeWeb()):
            result = manager.open_item(2, "火影-001", "view")
        self.assertTrue(result["ok"])
        self.assertEqual(opened, ["https://item.taobao.com/item.htm?id=1086638256748"])
        with self.assertRaises(RuntimeError):
            manager.open_item(2, "火影-001", "share")

    def test_clear_row_state_resets_execution(self):
        from desktop.jobs import MANAGER

        MANAGER._restored = True
        MANAGER.status = "idle"
        MANAGER.rows = [
            {
                "row": 2,
                "product_id": "晨光-001",
                "validation": "通过",
                "execution": "失败",
                "notice": "上传失败",
                "errors": ["上传失败"],
                "taobao_item_id": "1087608952245",
                "view_url": "https://item.taobao.com/item.htm?id=1087608952245",
                "edit_url": "https://item.taobao.com/item.htm?id=1087608952245&edit=1",
                "flow_version": "1",
                "flow_stage": "created",
                "run_status": "failed",
                "material_result": {"ok": False},
                "last_error": "上传失败",
                "sku_image_strategy": "slim_material",
            }
        ]
        MANAGER.products = [
            {
                "product": {"product_id": "晨光-001", "validation": "通过"},
                "meta": {
                    "row": 2,
                    "product_id": "晨光-001",
                    "validation": "通过",
                    "execution": "失败",
                    "taobao_item_id": "1087608952245",
                    "flow_version": "1",
                    "flow_stage": "created",
                    "last_error": "上传失败",
                },
            }
        ]
        result = MANAGER.clear_row_state(2, "晨光-001")
        self.assertTrue(result["ok"])
        row = MANAGER.rows[0]
        self.assertEqual(row["execution"], "未执行")
        self.assertEqual(row["taobao_item_id"], "")
        self.assertEqual(row["flow_stage"], "")
        self.assertEqual(row["errors"], [])
        meta = MANAGER.products[0]["meta"]
        self.assertEqual(meta["execution"], "未执行")
        self.assertEqual(meta["taobao_item_id"], "")
        self.assertEqual(meta["flow_stage"], "")

    def test_clear_row_state_running_rejects(self):
        from desktop.jobs import MANAGER

        MANAGER._restored = True
        MANAGER.status = "running"
        MANAGER.rows = [{"row": 2, "product_id": "晨光-001", "execution": "失败"}]
        MANAGER.products = []
        with self.assertRaises(RuntimeError):
            MANAGER.clear_row_state(2, "晨光-001")
        MANAGER.status = "idle"

    def test_item_clear_api(self):
        from fastapi.testclient import TestClient
        from desktop.server import app
        from desktop.jobs import MANAGER

        MANAGER._restored = True
        MANAGER.status = "idle"
        MANAGER.rows = [
            {
                "row": 3,
                "product_id": "火影-002",
                "validation": "通过",
                "execution": "已填写未提交",
                "taobao_item_id": "1084892550240",
                "flow_stage": "filled",
            }
        ]
        MANAGER.products = []
        client = TestClient(app)
        response = client.post("/api/item/clear", json={"row": 3, "product_id": "火影-002"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(MANAGER.rows[0]["execution"], "未执行")
        self.assertEqual(MANAGER.rows[0]["taobao_item_id"], "")
        missing = client.post("/api/item/clear", json={"row": 99, "product_id": "不存在"})
        self.assertEqual(missing.status_code, 400, missing.text)

    def test_item_open_api(self):
        from fastapi.testclient import TestClient
        from desktop.server import app
        from desktop.jobs import MANAGER

        MANAGER._restored = True
        MANAGER.rows = [{"row": 2, "product_id": "火影-001", "notice": "商品ID: 1086638256748"}]

        class FakeWeb:
            def open_item_url(self, url):
                return {"ok": True, "url": url, "via": "default-browser"}

        with unittest.mock.patch("desktop.jobs.load_web", return_value=FakeWeb()):
            client = TestClient(app)
            response = client.post("/api/item/open", json={"row": 2, "product_id": "火影-001", "action": "edit"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("itemId=1086638256748", response.json()["url"])

    def test_open_workspace_api(self):
        from fastapi.testclient import TestClient
        from desktop.server import app
        from desktop.jobs import MANAGER
        from test_workspace import write_mini_pack

        MANAGER._restored = True
        MANAGER.restored = False
        MANAGER.rows = []
        MANAGER.products = []
        MANAGER.workbook_path = ""
        MANAGER.workspace_path = ""
        root = Path(self.temp.name) / "ws"
        write_mini_pack(root / "火影-001", "规格A")
        write_mini_pack(root / "火影-002", "规格B")
        client = TestClient(app)
        response = client.post("/api/workspace/open", json={"path": str(root)})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data["count"], 2)
        self.assertTrue(data["workspace_path"].endswith("ws") or "ws" in data["workspace_path"])
        self.assertGreaterEqual(data["valid"], 0)
        defaults = client.get("/api/workspace/defaults")
        self.assertEqual(defaults.status_code, 200, defaults.text)
        saved = client.put("/api/workspace/defaults", json={"brand": "卡游", "price": 8.8, "stock": 12})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()["defaults"]["brand"], "卡游")
        rescan = client.post("/api/workspace/rescan")
        self.assertEqual(rescan.status_code, 200, rescan.text)
        self.assertEqual(rescan.json()["count"], 2)
        start = client.post("/api/job/start", json={"retry_failed": True})
        self.assertEqual(start.status_code, 400)

    def test_open_chrome_hides_when_logged_in(self):
        from desktop.jobs import JobManager
        chrome = Path(self.temp.name) / "chrome.exe"
        chrome.write_bytes(b"mz")
        hidden = []
        revealed = []

        class Settings:
            chrome_path = str(chrome)
            chrome_profile = self.temp.name
            cdp_port = 9222

        class FakeWeb:
            def reload_paths(self):
                pass

            def detect_chrome(self):
                return chrome

            def cdp_available(self):
                return True

            def cdp_version(self):
                return {"Browser": "Chrome/1"}

            def login_status_from_tabs(self):
                return {"logged_in": True, "blocker": "", "url": "https://myseller.taobao.com/home.htm"}

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                return "hidden"

            def hide_automation_chrome(self, **kwargs):
                hidden.append(True)
                return True

            def prune_automation_tabs(self, **kwargs):
                return {"closed": 0, "kept": 1}

            def reveal_automation_chrome(self, **kwargs):
                revealed.append(True)
                return True

            def reveal_seller_chrome(self):
                revealed.append(True)

        manager = JobManager()
        manager._restored = True
        with unittest.mock.patch("desktop.jobs.paths.configure_environ", return_value=Settings), \
             unittest.mock.patch("desktop.jobs.paths.load_settings", return_value=Settings), \
             unittest.mock.patch("desktop.jobs.load_web", return_value=FakeWeb()):
            with unittest.mock.patch("desktop.jobs.time.monotonic", side_effect=[100.0, 103.0]):
                manager.chrome_status()
                status = manager.open_chrome()
        self.assertTrue(hidden)
        self.assertFalse(revealed)
        self.assertIn("已收起", status["notice"])
        self.assertNotIn("前台", status["notice"])

    def test_resume_hides_browser_without_waiting_for_login_settle(self):
        from desktop.jobs import JobManager
        web = unittest.mock.Mock()
        web.cdp_available.return_value = True
        manager = JobManager()
        with unittest.mock.patch("desktop.jobs.paths.load_settings", return_value=unittest.mock.Mock(debug_browser=False)), \
             unittest.mock.patch("desktop.jobs.load_web", return_value=web):
            manager._hide_chrome_for_fill()
        web.hide_automation_chrome.assert_called_once()
        web.login_status_from_tabs.assert_not_called()
        web.reveal_automation_chrome.assert_not_called()

    def test_job_cleanup_preserves_failed_paused_and_stopped_browser(self):
        from desktop.jobs import JobManager
        for status, execution in [("paused", "暂停"), ("error", "失败"),
                                  ("stopped", ""), ("done", "失败"),
                                  ("done", "结果待核实"), ("done", "已入库")]:
            with self.subTest(status=status, execution=execution):
                manager = JobManager()
                manager.status = status
                manager.rows = [{"execution": execution}]
                manager._hide_chrome_after_login = True
                web = unittest.mock.Mock()
                with unittest.mock.patch("desktop.jobs.load_web", return_value=web), \
                     unittest.mock.patch.object(manager, "_reveal_chrome_for_login") as reveal, \
                     unittest.mock.patch.object(manager, "_hide_chrome_for_fill") as hide:
                    manager._finish_browser()
                if status == "done" and execution == "已入库":
                    web.prune_automation_tabs.assert_called_once()
                    hide.assert_called_once()
                    reveal.assert_not_called()
                else:
                    web.prune_automation_tabs.assert_not_called()
                    hide.assert_not_called()
                    reveal.assert_called_once()
                    self.assertFalse(manager._hide_chrome_after_login)

    def test_login_status_requires_stable_seller_url(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager._restored = True
        login = {"logged_in": True, "blocker": "", "url": "https://myseller.taobao.com/home.htm"}
        with unittest.mock.patch("desktop.jobs.time.monotonic", side_effect=[10.0, 11.0, 13.1]):
            first = manager._stable_login_status(login)
            second = manager._stable_login_status(login)
            stable = manager._stable_login_status(login)
        self.assertFalse(first["logged_in"])
        self.assertFalse(second["logged_in"])
        self.assertEqual(second["blocker"], "正在确认登录状态")
        self.assertTrue(stable["logged_in"])

        reset = manager._stable_login_status(
            {"logged_in": False, "blocker": "登录页", "url": "https://loginmyseller.taobao.com/"}
        )
        self.assertFalse(reset["logged_in"])
        self.assertEqual(manager._login_candidate_url, "")

    def test_login_status_stays_confirmed_across_url_changes(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager._restored = True
        url_a = {"logged_in": True, "blocker": "", "url": "https://myseller.taobao.com/home.htm"}
        url_b = {"logged_in": True, "blocker": "", "url": "https://myseller.taobao.com/item.htm?id=1"}
        clock = iter([10.0, 13.5])
        with unittest.mock.patch("desktop.jobs.time.monotonic", lambda: next(clock)):
            settling = manager._stable_login_status(url_a)
            confirmed = manager._stable_login_status(url_a)
            # 已确认后 URL 变化（任务运行期间页面跳转）不翻转登录状态
            after_jump = manager._stable_login_status(url_b)
        self.assertEqual(settling["blocker"], "正在确认登录状态")
        self.assertTrue(confirmed["logged_in"])
        self.assertTrue(after_jump["logged_in"])

        # 出现登录页解除确认后，需重新经过稳定窗口
        manager._stable_login_status({"logged_in": False, "blocker": "登录页", "url": "https://login.taobao.com/"})
        clock = iter([20.0, 21.0, 24.0])
        with unittest.mock.patch("desktop.jobs.time.monotonic", lambda: next(clock)):
            reconfirming = manager._stable_login_status(url_a)
            still_settling = manager._stable_login_status(url_a)
            reconfirmed = manager._stable_login_status(url_a)
        self.assertEqual(reconfirming["blocker"], "正在确认登录状态")
        self.assertEqual(still_settling["blocker"], "正在确认登录状态")
        self.assertTrue(reconfirmed["logged_in"])

    def _auto_check_manager(self, web, timeout=0.5, settle=0.0, grace=0.0):
        from desktop.jobs import JobManager

        chrome = Path(self.temp.name) / "chrome.exe"
        chrome.write_bytes(b"mz")

        class Settings:
            chrome_path = str(chrome)
            chrome_profile = self.temp.name
            cdp_port = 9222
            debug_browser = False

        manager = JobManager()
        manager._restored = True
        manager._login_check_timeout = timeout
        manager._login_check_settle = settle
        manager._login_check_grace = grace
        patches = [
            unittest.mock.patch("desktop.jobs.paths.configure_environ", return_value=Settings),
            unittest.mock.patch("desktop.jobs.paths.load_settings", return_value=Settings),
            unittest.mock.patch("desktop.jobs.paths.browser_exe_path", return_value=chrome),
            unittest.mock.patch("desktop.jobs.load_web", return_value=web),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)
        return manager

    def test_auto_login_check_confirms_connected_session_without_click(self):
        seller_tab = {"id": "t1", "type": "page", "url": "https://myseller.taobao.com/home.htm"}
        spawned = []

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return True

            def cdp_version(self):
                return {"Browser": "Chrome/1"}

            def cdp_tabs(self):
                return [seller_tab]

            def login_status_from_tabs(self, tabs=None):
                return {"logged_in": True, "blocker": "", "url": seller_tab["url"]}

            def is_login_url(self, url):
                return "login" in str(url)

            def is_seller_url(self, url):
                return "myseller" in str(url)

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                spawned.append((focus, hide_if_logged_in))
                return "started"

            def activate_seller_tab(self, tabs=None):
                return seller_tab

        manager = self._auto_check_manager(FakeWeb())
        manager._run_login_check()
        self.assertEqual(spawned, [], "CDP 已连接时不应重复启动浏览器")
        self.assertTrue(manager._login_confirmed)
        status = manager.chrome_status()
        self.assertTrue(status["logged_in"])
        self.assertFalse(status["checking"])
        self.assertTrue(any("无需再次登录" in (entry.get("message") or "") for entry in manager.logs))

    def test_auto_login_check_starts_hidden_browser_and_opens_seller_home(self):
        state = {"cdp": False, "tabs": []}
        spawned = []

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return state["cdp"]

            def cdp_version(self):
                return {"Browser": "Chrome/1"}

            def cdp_tabs(self):
                return list(state["tabs"])

            def login_status_from_tabs(self, tabs=None):
                tabs = state["tabs"]
                urls = [str(tab.get("url") or "") for tab in tabs]
                logins = [u for u in urls if "login" in u]
                sellers = [u for u in urls if "myseller" in u]
                if logins:
                    return {"logged_in": False, "blocker": "登录页", "url": logins[0]}
                if sellers:
                    return {"logged_in": True, "blocker": "", "url": sellers[0]}
                return {"logged_in": False, "blocker": "未打开卖家中心", "url": urls[0] if urls else ""}

            def is_login_url(self, url):
                return "login" in str(url)

            def is_seller_url(self, url):
                return "myseller" in str(url)

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                spawned.append((focus, hide_if_logged_in))
                state["cdp"] = True
                return "started"

            def activate_seller_tab(self, tabs=None):
                state["tabs"].append({"id": "t2", "type": "page", "url": "https://myseller.taobao.com/"})
                return None

        manager = self._auto_check_manager(FakeWeb())
        with unittest.mock.patch("desktop.jobs.time.sleep", lambda _seconds: None):
            manager._run_login_check()
        self.assertEqual(spawned, [(False, True)], "应后台静默启动（不聚焦、默认收起）")
        self.assertTrue(manager._login_confirmed)
        status = manager.chrome_status()
        self.assertTrue(status["logged_in"])
        self.assertFalse(status["checking"])

    def test_auto_login_check_reports_login_page_without_confirming(self):
        state = {"cdp": False, "tabs": [{"id": "t1", "type": "page", "url": "https://loginmyseller.taobao.com/?redirect=x"}]}
        spawned = []

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return state["cdp"]

            def cdp_version(self):
                return {"Browser": "Chrome/1"}

            def cdp_tabs(self):
                return list(state["tabs"])

            def login_status_from_tabs(self, tabs=None):
                return {"logged_in": False, "blocker": "登录页", "url": state["tabs"][0]["url"]}

            def is_login_url(self, url):
                return "login" in str(url)

            def is_seller_url(self, url):
                return "myseller" in str(url)

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                spawned.append((focus, hide_if_logged_in))
                state["cdp"] = True
                return "started"

            def activate_seller_tab(self, tabs=None):
                raise AssertionError("已出现登录页时不应再补开卖家首页")

        manager = self._auto_check_manager(FakeWeb())
        with unittest.mock.patch("desktop.jobs.time.sleep", lambda _seconds: None):
            manager._run_login_check()
        self.assertEqual(spawned, [(False, True)])
        self.assertFalse(manager._login_confirmed)
        status = manager.chrome_status()
        self.assertFalse(status["logged_in"])
        self.assertEqual(status["blocker"], "登录页")

    def test_auto_login_check_skipped_while_job_running(self):
        spawned = []

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return False

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                spawned.append(True)
                return "started"

        manager = self._auto_check_manager(FakeWeb())
        manager.status = "running"
        manager._run_login_check()
        self.assertEqual(spawned, [])
        self.assertEqual(manager._auto_check_state, "done")

    def test_auto_login_check_skipped_without_browser(self):
        spawned = []

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return False

            def start_persistent_chrome(self, focus=False, hide_if_logged_in=True):
                spawned.append(True)
                return "started"

        manager = self._auto_check_manager(FakeWeb())
        with unittest.mock.patch("desktop.jobs.paths.browser_exe_path", return_value=None):
            manager._run_login_check()
        self.assertEqual(spawned, [])
        self.assertEqual(manager._auto_check_state, "done")
        self.assertTrue(any("未找到可用浏览器" in (entry.get("message") or "") for entry in manager.logs))

    def test_chrome_status_reports_checking_while_auto_check_runs(self):
        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return False

        manager = self._auto_check_manager(FakeWeb())
        for state, expected in (("running", True), ("done", False)):
            with self.subTest(state=state):
                manager._auto_check_state = state
                self.assertEqual(manager.chrome_status()["checking"], expected)

    def test_begin_login_check_runs_once_per_process(self):
        state = {"cdp": True, "tabs": [{"id": "t1", "type": "page", "url": "https://myseller.taobao.com/home.htm"}]}

        class FakeWeb:
            def reload_paths(self):
                pass

            def cdp_available(self):
                return state["cdp"]

            def cdp_version(self):
                return {"Browser": "Chrome/1"}

            def cdp_tabs(self):
                return list(state["tabs"])

            def login_status_from_tabs(self, tabs=None):
                return {"logged_in": True, "blocker": "", "url": state["tabs"][0]["url"]}

            def is_login_url(self, url):
                return "login" in str(url)

            def is_seller_url(self, url):
                return "myseller" in str(url)

            def activate_seller_tab(self, tabs=None):
                return None

        manager = self._auto_check_manager(FakeWeb())
        self.assertTrue(manager.begin_login_check())
        self.assertFalse(manager.begin_login_check(), "检测期间重复触发应被忽略")
        deadline = time.time() + 5
        while manager._auto_check_state != "done" and time.time() < deadline:
            time.sleep(0.02)
        self.assertEqual(manager._auto_check_state, "done")
        self.assertFalse(manager.begin_login_check(), "每进程只自动检测一次")
        self.assertTrue(manager._login_confirmed)

    def test_check_login_endpoint_triggers_manager_once(self):
        from fastapi.testclient import TestClient

        from desktop import jobs as jobs_module
        from desktop.server import app

        client = TestClient(app)
        with unittest.mock.patch.object(jobs_module.MANAGER, "begin_login_check", return_value=True) as begin:
            response = client.post("/api/chrome/check-login")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {"started": True})
        begin.assert_called_once_with()

    def test_watch_progress_does_not_leak_end_across_items(self):
        from desktop.jobs import JobManager
        manager = JobManager()
        manager._restored = True
        manager.current_id = "火影-002"
        watch = {"last_id": "火影-001", "last_label": "提交"}
        leftover = [{"step": "end", "execution": "提交失败", "notice": "错误(2) 帮助 反馈"}]
        manager._consume_fill_progress(leftover, watch)
        self.assertFalse(any("本条结束" in (entry.get("message") or "") for entry in manager.logs))
        self.assertEqual(watch["last_id"], "火影-002")
        manager._consume_fill_progress(
            [{"step": "tab", "url": "https://item.upload.taobao.com/sell/v2/publish.htm"}],
            watch,
        )
        messages = [entry.get("message") or "" for entry in manager.logs]
        self.assertTrue(any("定位填写页" in msg for msg in messages))
        self.assertFalse(any("本条结束" in msg for msg in messages))


class StallWatchdogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        # 保存并恢复 QIANNIU_APPDATA：临时目录在 cleanup 后会被删除，
        # 泄漏会污染后续测试的会话持久化读取。
        self._old_appdata = os.environ.get("QIANNIU_APPDATA")
        os.environ["QIANNIU_APPDATA"] = self.temp.name
        from desktop.paths import configure_environ
        configure_environ()

    def tearDown(self):
        os.environ.pop("QIANNIU_STALL_SECONDS", None)
        if self._old_appdata is None:
            os.environ.pop("QIANNIU_APPDATA", None)
        else:
            os.environ["QIANNIU_APPDATA"] = self._old_appdata

    def test_stall_seconds_env_override(self):
        from desktop import jobs

        os.environ["QIANNIU_STALL_SECONDS"] = "90"
        self.assertEqual(jobs.stall_seconds(), 90)
        os.environ["QIANNIU_STALL_SECONDS"] = "not-a-number"
        self.assertEqual(jobs.stall_seconds(), jobs.DEFAULT_STALL_SECONDS)
        os.environ["QIANNIU_STALL_SECONDS"] = "10"
        self.assertEqual(jobs.stall_seconds(), 60, "阈值下限为 60 秒")

    def test_watchdog_marks_running_task_stalled(self):
        import time as time_module

        from desktop.jobs import JobManager

        manager = JobManager()
        manager.status = "running"
        manager.activity_at = time_module.time() - 10 ** 6
        manager._reveal_chrome_for_login = unittest.mock.Mock()
        manager._check_stalled()
        self.assertTrue(manager.stalled)
        self.assertIn("无任何进展", manager.blocker)
        self.assertTrue(any("无任何进展" in (entry.get("message") or "") for entry in manager.logs))
        manager._reveal_chrome_for_login.assert_called_once()

    def test_watchdog_ignores_non_running_task(self):
        import time as time_module

        from desktop.jobs import JobManager

        manager = JobManager()
        manager.status = "idle"
        manager.activity_at = time_module.time() - 10 ** 6
        manager._check_stalled()
        self.assertFalse(manager.stalled)
        self.assertEqual(manager.blocker, "")

    def test_activity_touch_recovers_from_stall(self):
        import time as time_module

        from desktop.jobs import JobManager

        manager = JobManager()
        manager.status = "running"
        manager.stalled = True
        manager.activity_at = time_module.time() - 10 ** 6
        manager._touch_activity()
        self.assertFalse(manager.stalled)
        self.assertTrue(any("恢复进展" in (entry.get("message") or "") for entry in manager.logs))

    def test_watchdog_state_resets_on_snapshot_payload(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager.stalled = True
        self.assertTrue(manager.snapshot()["stalled"])
        manager.stalled = False
        self.assertFalse(manager.snapshot()["stalled"])


class PauseReasonCodeTests(unittest.TestCase):
    """分类行为直接测 web_fill.pipeline；jobs 的 helper 只用 mock 验证转发。

    不能在测试运行期真实调用 load_seller().load_web_fill()：它会卸载并重载
    sys.modules 里的 web_fill*，破坏其他测试持有旧模块对象的 patch。
    """

    @staticmethod
    def _classifier():
        from web_fill.pipeline import pause_reason_code

        return pause_reason_code

    def test_explicit_code_wins(self):
        classify = self._classifier()
        self.assertEqual(classify("PAUSE:captcha:已有商品核验遇到安全验证"), "captcha")
        self.assertEqual(classify("PAUSE:login:需要登录卖家中心"), "login")
        self.assertEqual(classify("PAUSE:slider:请完成滑块验证"), "slider")

    def test_keyword_fallback(self):
        classify = self._classifier()
        self.assertEqual(classify("PAUSE:检测到滑块验证，请人工处理"), "slider")
        self.assertEqual(classify("Error: 页面跳转到登录页"), "login")
        self.assertEqual(classify("淘宝触发“滑块验证”，请手动完成"), "slider")

    def test_non_hard_pause_returns_empty(self):
        classify = self._classifier()
        self.assertEqual(classify("PAUSE:规格未写入 颜色"), "")
        self.assertEqual(classify("PAUSE:用户已暂停；保留当前发布页"), "")
        self.assertEqual(classify("普通失败: 标题超长"), "")
        self.assertEqual(classify(""), "")

    def test_jobs_helper_forwards_to_classifier(self):
        from types import SimpleNamespace

        from desktop import jobs

        seller = SimpleNamespace(load_web_fill=lambda: SimpleNamespace(
            pause_reason_code=lambda text: "captcha" if "安全验证" in text else ""))
        with unittest.mock.patch.object(jobs, "load_seller", return_value=seller):
            self.assertEqual(jobs.pause_reason_code("PAUSE:captcha:已有商品核验遇到安全验证"), "captcha")
            self.assertEqual(jobs.pause_reason_code("PAUSE:规格未写入"), "")

    def test_jobs_helper_falls_back_to_keywords(self):
        from desktop import jobs

        with unittest.mock.patch.object(jobs, "load_seller", side_effect=RuntimeError("不可用")):
            self.assertEqual(jobs.pause_reason_code("PAUSE:检测到滑块验证"), "slider")
            self.assertEqual(jobs.pause_reason_code("跳转到登录页"), "login")
            self.assertEqual(jobs.pause_reason_code("普通失败"), "")


class JobTimingTests(unittest.TestCase):
    def test_serialize_row_includes_duration_seconds(self):
        from desktop.jobs import serialize_row

        row = serialize_row({"row": 2, "product_id": "A", "duration_seconds": 205.3, "product": {}})
        self.assertEqual(row["duration_seconds"], 205.3)
        self.assertIsNone(serialize_row({"row": 3, "product_id": "B", "product": {}})["duration_seconds"])

    def test_format_duration(self):
        from desktop.jobs import format_duration

        self.assertEqual(format_duration(45), "45秒")
        self.assertEqual(format_duration(205.4), "3分25秒")
        self.assertEqual(format_duration(5000), "1小时23分")
        self.assertEqual(format_duration(0), "")
        self.assertEqual(format_duration(None), "")

    def test_snapshot_reports_job_timing(self):
        from desktop.jobs import JobManager

        manager = JobManager()
        manager.workbook_path = "x.xlsx"
        manager.status = "running"
        manager.started_at = time.time() - 65.0
        manager.item_durations = [120.0, 90.5]
        snap = manager.snapshot()
        self.assertAlmostEqual(snap["elapsed_seconds"], 65.0, delta=3.0)
        self.assertEqual(snap["avg_item_seconds"], 105.2)
        self.assertIsNone(snap["finished_at"])

        manager.status = "done"
        manager.finished_at = manager.started_at + 240.0
        snap = manager.snapshot()
        self.assertAlmostEqual(snap["elapsed_seconds"], 240.0, delta=0.5)
        self.assertAlmostEqual(snap["finished_at"], manager.started_at + 240.0)

        manager.started_at = 0.0
        manager.finished_at = 0.0
        manager.item_durations = []
        self.assertIsNone(manager.snapshot()["elapsed_seconds"])
        self.assertIsNone(manager.snapshot()["avg_item_seconds"])

    def test_job_history_parses_result_journals(self):
        from types import SimpleNamespace

        from desktop import jobs

        with tempfile.TemporaryDirectory() as directory:
            results = Path(directory) / "results"
            results.mkdir(parents=True)
            now = time.time()
            new_file = results / "清单A.结果-20260930-120000-000000.json"
            new_file.write_text(json.dumps({
                "source": "G:/x/清单A.xlsx",
                "started_at": now - 600,
                "finished_at": now - 60,
                "duration_seconds": 540.0,
                "results": [
                    {"product_id": "A", "taobao_item_id": "111", "execution": "已入库，图片已核验", "duration_seconds": 300.0},
                    {"product_id": "B", "taobao_item_id": "", "execution": "失败", "duration_seconds": 240.0},
                ],
            }, ensure_ascii=False), encoding="utf-8")
            old_file = results / "清单B.结果-20260929-100000-000000.json"
            old_file.write_text(json.dumps({
                "source": "G:/x/清单B.xlsx",
                "results": [{"product_id": "C", "taobao_item_id": "333", "execution": "已填写未提交"}],
            }, ensure_ascii=False), encoding="utf-8")
            (results / "坏文件.结果-20260928-100000-000000.json").write_text("{broken", encoding="utf-8")
            (results / "settings.json").write_text("{}", encoding="utf-8")
            # job_history 按 mtime 排序；连续写入可能落在同一时间刻度导致顺序
            # 不稳定，这里显式固定新旧文件的修改时间。
            os.utime(new_file, (now, now))
            os.utime(old_file, (now - 86400, now - 86400))

            with unittest.mock.patch.object(
                jobs.paths, "load_settings", return_value=SimpleNamespace(results_dir=str(results))
            ):
                items = jobs.MANAGER.job_history()

            self.assertEqual([item["file"] for item in items], [new_file.name, old_file.name])
            newest = items[0]
            self.assertEqual(newest["total"], 2)
            self.assertEqual(newest["succeeded"], 1)
            self.assertEqual(newest["failed"], 1)
            self.assertAlmostEqual(newest["avg_seconds"], 270.0)
            self.assertAlmostEqual(newest["duration_seconds"], 540.0)
            oldest = items[1]
            self.assertEqual(oldest["total"], 1)
            self.assertEqual(oldest["succeeded"], 1)
            self.assertEqual(oldest["failed"], 0)
            self.assertIsNone(oldest["duration_seconds"])
            self.assertIsNone(oldest["avg_seconds"])
            self.assertIsNone(oldest["started_at"])


if __name__ == "__main__":
    unittest.main()
