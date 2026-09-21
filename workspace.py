"""工作空间：一款一图片夹 + 一张商品清单.xlsx。CLI 与桌面共用。"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from 商品解析 import IMAGE_EXTS, load_registry, scan_image_pack

DEFAULTS_NAME = "批次默认.json"
WORKBOOK_NAME = "商品清单.xlsx"
MISSING_STATUS = "资料缺失"
MISSING_NOTICE = "图片夹不存在或没有可识别图片"
TITLE_MAX_WIDTH = 60

DEFAULT_DEFAULTS = {
    "品牌": "",
    "商品属性模板": "中性笔",
    "物流模板": "48小时",
    "销售模板": "仓库多规格",
    "价格": 9.9,
    "库存": 20,
}

_API_TO_FILE = {
    "brand": "品牌",
    "attributes_template": "商品属性模板",
    "logistics_template": "物流模板",
    "sales_template": "销售模板",
    "price": "价格",
    "stock": "库存",
}
_FILE_TO_API = {value: key for key, value in _API_TO_FILE.items()}


def _present(value):
    return value is not None and str(value).strip() != ""


def _seller():
    existing = sys.modules.get("千牛自动上架")
    if existing is not None:
        return existing
    import importlib.util

    path = Path(__file__).with_name("千牛自动上架.py")
    spec = importlib.util.spec_from_file_location("千牛自动上架", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules["千牛自动上架"] = module
    spec.loader.exec_module(module)
    return module


def title_width(text):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text or ""))


def clip_title(text, max_width=TITLE_MAX_WIDTH):
    out = []
    width = 0
    for char in str(text or ""):
        step = 2 if ord(char) > 127 else 1
        if width + step > max_width:
            break
        out.append(char)
        width += step
    return "".join(out)


def template_registry():
    try:
        data = load_registry()
    except Exception:
        data = {}
    return {
        "attributes": list(data.get("attributes") or ["中性笔"]),
        "logistics": list(data.get("logistics") or ["48小时", "24小时"]),
        "sales": list(data.get("sales") or ["仓库多规格"]),
    }


def normalize_defaults(data=None):
    merged = dict(DEFAULT_DEFAULTS)
    raw = dict(data or {})
    for api_key, file_key in _API_TO_FILE.items():
        if _present(raw.get(file_key)):
            merged[file_key] = raw.get(file_key)
        elif _present(raw.get(api_key)):
            merged[file_key] = raw.get(api_key)
    price = merged.get("价格")
    stock = merged.get("库存")
    try:
        merged["价格"] = float(price) if price not in (None, "") else DEFAULT_DEFAULTS["价格"]
    except (TypeError, ValueError):
        merged["价格"] = DEFAULT_DEFAULTS["价格"]
    try:
        merged["库存"] = int(stock) if stock not in (None, "") else DEFAULT_DEFAULTS["库存"]
    except (TypeError, ValueError):
        merged["库存"] = DEFAULT_DEFAULTS["库存"]
    for key in ("品牌", "商品属性模板", "物流模板", "销售模板"):
        merged[key] = str(merged.get(key) or "").strip()
    if not merged["商品属性模板"]:
        merged["商品属性模板"] = DEFAULT_DEFAULTS["商品属性模板"]
    if not merged["物流模板"]:
        merged["物流模板"] = DEFAULT_DEFAULTS["物流模板"]
    if not merged["销售模板"]:
        merged["销售模板"] = DEFAULT_DEFAULTS["销售模板"]
    return merged


def defaults_for_api(data=None):
    merged = normalize_defaults(data)
    return {api_key: merged.get(file_key) for api_key, file_key in _API_TO_FILE.items()}


def workbook_path(root):
    return Path(root).expanduser().resolve() / WORKBOOK_NAME


def defaults_path(root):
    return Path(root).expanduser().resolve() / DEFAULTS_NAME


def load_defaults(root):
    path = defaults_path(root)
    data = {}
    if path.is_file():
        try:
            loaded = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                data = loaded
        except (OSError, json.JSONDecodeError):
            data = {}
    return normalize_defaults(data)


def save_defaults(root, data=None):
    folder = Path(root).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    current = load_defaults(folder)
    raw = dict(data or {})
    patch = {}
    for api_key, file_key in _API_TO_FILE.items():
        if api_key in raw and raw[api_key] is not None:
            patch[file_key] = raw[api_key]
        elif file_key in raw and raw[file_key] is not None:
            patch[file_key] = raw[file_key]
    merged = normalize_defaults({**current, **patch})
    path = defaults_path(folder)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)
    return merged


def _relative_folder(root, folder):
    try:
        return Path(os.path.relpath(folder, root)).as_posix()
    except ValueError:
        return str(folder)


def _pack_has_media(pack):
    return bool(
        pack.get("main_1_1")
        or pack.get("main_3_4")
        or pack.get("details")
        or pack.get("skus")
    )


def _sku_summary(pack):
    items = []
    for sku in pack.get("skus") or []:
        items.append({
            "slot": str(sku.get("slot") or "").strip(),
            "name": str(sku.get("name") or "").strip(),
        })
    return [item for item in items if item["slot"]]


def scan_workspace(root):
    """扫描工作空间一级子目录，忽略夹内 Excel。"""
    folder = Path(root).expanduser().resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"工作空间不存在或不是目录: {folder}")
    found = []
    skipped = []
    errors = []
    for child in sorted(folder.iterdir(), key=lambda item: item.name.lower()):
        if not child.is_dir():
            continue
        name = child.name
        if name.startswith("."):
            skipped.append({"folder": name, "reason": "隐藏目录"})
            continue
        files = [path for path in child.iterdir() if path.is_file()]
        images = [path for path in files if path.suffix.lower() in IMAGE_EXTS]
        if not files:
            skipped.append({"folder": name, "reason": "空目录"})
            continue
        if not images:
            errors.append({"folder": name, "error": "没有可识别图片"})
            continue
        try:
            pack = scan_image_pack(child)
        except Exception as exc:
            errors.append({"folder": name, "error": str(exc)})
            continue
        if not _pack_has_media(pack):
            errors.append({"folder": name, "error": "没有可识别图片"})
            continue
        found.append({
            "product_id": name,
            "name": name,
            "path": str(child),
            "relative": _relative_folder(folder, child),
            "main_count": len(pack.get("main_1_1") or []),
            "portrait_count": len(pack.get("main_3_4") or []),
            "detail_count": len(pack.get("details") or []),
            "sku_count": len(pack.get("skus") or []),
            "skus": _sku_summary(pack),
        })
    return {
        "root": str(folder),
        "folders": found,
        "skipped": skipped,
        "errors": errors,
    }


def _read_sheet_dicts(ws, headers_wanted=None):
    rows = ws.iter_rows(values_only=True)
    headers = [str(cell or "").strip() for cell in next(rows, ())]
    items = []
    for values in rows:
        data = dict(zip(headers, values))
        if headers_wanted:
            if not any(_present(data.get(key)) for key in headers_wanted):
                continue
        elif not any(_present(value) for value in data.values()):
            continue
        items.append(data)
    return headers, items


def _read_existing(path):
    from openpyxl import load_workbook

    seller = _seller()
    wb = load_workbook(path)
    try:
        sheet = wb["商品清单"] if "商品清单" in wb.sheetnames else wb.worksheets[0]
        _, products = _read_sheet_dicts(sheet)
        sku_rows = []
        if "SKU规格" in wb.sheetnames:
            _, sku_rows = _read_sheet_dicts(wb["SKU规格"], seller.SKU_HEADERS)
        log_rows = []
        if "处理日志" in wb.sheetnames:
            log = wb["处理日志"]
            for row in log.iter_rows(min_row=2, values_only=True):
                if any(_present(value) for value in row):
                    log_rows.append(list(row))
        return products, sku_rows, log_rows
    finally:
        wb.close()


def _as_template_row(data):
    seller = _seller()
    row = {field: data.get(field, "") for field in seller.TEMPLATE_HEADERS}
    if not _present(row.get("品牌*")) and _present(data.get("品牌")):
        row["品牌*"] = data.get("品牌")
    return row


def _new_product_row(folder, defaults):
    seller = _seller()
    name = folder["product_id"]
    brand = str(defaults.get("品牌") or "").strip()
    raw_title = name
    if brand and brand not in name:
        raw_title = f"{brand}{name}"
    title = clip_title(raw_title)
    model = name if title != name else ""
    row = {field: "" for field in seller.TEMPLATE_HEADERS}
    row.update({
        "商品标识*": name,
        "商品标题*": title,
        "品牌*": brand,
        "型号": model,
        "图片包路径*": folder["relative"],
        "商品属性模板*": defaults.get("商品属性模板") or "中性笔",
        "物流模板*": defaults.get("物流模板") or "48小时",
        "销售模板*": defaults.get("销售模板") or "仓库多规格",
        "价格*": defaults.get("价格"),
        "库存*": defaults.get("库存"),
        "处理状态": "待处理",
        "错误信息": "",
    })
    return row


def _clear_missing(row):
    if str(row.get("处理状态") or "").strip() == MISSING_STATUS:
        row["处理状态"] = "待处理"
        notice = str(row.get("错误信息") or "")
        if MISSING_NOTICE in notice or "图片夹" in notice:
            row["错误信息"] = ""
    return row


def _mark_missing(row):
    updated = dict(row)
    updated["处理状态"] = MISSING_STATUS
    updated["错误信息"] = MISSING_NOTICE
    return updated


def _sku_key(row):
    return (str(row.get("商品标识") or "").strip(), str(row.get("色号") or "").strip())


def _merge_rows(existing_products, existing_skus, scan, defaults):
    seller = _seller()
    folder_by_id = {item["product_id"]: item for item in scan.get("folders") or []}
    products = []
    seen = set()
    for raw in existing_products or []:
        row = _as_template_row(raw)
        product_id = str(row.get("商品标识*") or "").strip()
        if not product_id:
            continue
        seen.add(product_id)
        folder = folder_by_id.get(product_id)
        if folder:
            row["图片包路径*"] = folder["relative"]
            _clear_missing(row)
        else:
            row = _mark_missing(row)
        products.append(row)
    added = []
    for folder in scan.get("folders") or []:
        if folder["product_id"] in seen:
            continue
        products.append(_new_product_row(folder, defaults))
        added.append(folder["product_id"])
    sku_rows = []
    emitted = set()
    for raw in existing_skus or []:
        item = {field: raw.get(field, "") for field in seller.SKU_HEADERS}
        key = _sku_key(item)
        if not key[1] or key in emitted:
            continue
        sku_rows.append(item)
        emitted.add(key)
    for folder in scan.get("folders") or []:
        for sku in folder.get("skus") or []:
            key = (folder["product_id"], sku["slot"])
            if not key[1] or key in emitted:
                continue
            sku_rows.append({
                "商品标识": folder["product_id"],
                "色号": sku["slot"],
                "规格名称": "",
                "价格": "",
                "库存": "",
            })
            emitted.add(key)
    missing = [
        str(row.get("商品标识*") or "")
        for row in products
        if str(row.get("处理状态") or "") == MISSING_STATUS
    ]
    return products, sku_rows, added, missing


def sync_workbook(root, defaults=None):
    """生成或按商品标识合并 商品清单.xlsx，返回路径和扫描摘要。"""
    folder = Path(root).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    scan = scan_workspace(folder)
    saved_defaults = save_defaults(folder, defaults if defaults is not None else load_defaults(folder))
    target = workbook_path(folder)
    existing_products, existing_skus, log_rows = [], [], []
    created = not target.is_file()
    if target.is_file():
        try:
            existing_products, existing_skus, log_rows = _read_existing(target)
        except Exception:
            existing_products, existing_skus, log_rows = [], [], []
            created = True
    products, sku_rows, added, missing = _merge_rows(
        existing_products, existing_skus, scan, saved_defaults
    )
    seller = _seller()
    seller.write_listing_workbook(target, products, sku_rows, log_rows)
    return {
        "path": str(target),
        "root": str(folder),
        "created": created,
        "added": added,
        "missing": missing,
        "defaults": saved_defaults,
        "scan": scan,
        "count": len(products),
    }
