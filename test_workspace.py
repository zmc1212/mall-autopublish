import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

import workspace
from workspace import MISSING_STATUS, clip_title, scan_workspace, sync_workbook


def write_mini_pack(folder, spec_name="规格A"):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (8, 8), "red").save(folder / "宝贝主图01.jpg")
    Image.new("RGB", (8, 8), "blue").save(folder / "详情01.jpg")
    Image.new("RGB", (8, 8), "green").save(folder / f"颜色01-{spec_name}.jpg")
    return folder


def list_naruto_packs():
    testdata = Path(__file__).resolve().parent / "testdata"
    if not testdata.is_dir():
        return []
    return sorted(
        [path for path in testdata.iterdir() if path.is_dir() and "火影" in path.name],
        key=lambda item: item.name,
    )


def link_folder(source, dest):
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        return dest
    source = Path(source).resolve()
    if os.name == "nt":
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(dest), str(source)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if result.returncode != 0 or not dest.exists():
            raise OSError(result.stderr or result.stdout or "mklink failed")
        return dest
    dest.symlink_to(source, target_is_directory=True)
    return dest


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def _titles(self, path):
        wb = load_workbook(path, data_only=True)
        try:
            sheet = wb["商品清单"]
            headers = [cell.value for cell in sheet[1]]
            rows = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(values):
                    continue
                rows.append(dict(zip(headers, values)))
            return rows
        finally:
            wb.close()

    def test_empty_workspace(self):
        scan = scan_workspace(self.root)
        self.assertEqual(scan["folders"], [])
        synced = sync_workbook(self.root)
        self.assertTrue(Path(synced["path"]).is_file())
        self.assertEqual(synced["count"], 0)
        self.assertTrue((self.root / "批次默认.json").is_file())

    def test_two_packs_merge_keeps_edited_title(self):
        write_mini_pack(self.root / "火影-001", "规格A")
        write_mini_pack(self.root / "火影-002", "规格B")
        synced = sync_workbook(self.root, {
            "品牌": "卡游",
            "商品属性模板": "中性笔",
            "物流模板": "48小时",
            "销售模板": "仓库多规格",
            "价格": 9.9,
            "库存": 20,
        })
        self.assertEqual(synced["count"], 2)
        path = Path(synced["path"])
        wb = load_workbook(path)
        sheet = wb["商品清单"]
        headers = [cell.value for cell in sheet[1]]
        title_col = headers.index("商品标题*") + 1
        id_col = headers.index("商品标识*") + 1
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, id_col).value == "火影-001":
                sheet.cell(row, title_col).value = "卡游手改标题不要覆盖"
        wb.save(path)
        wb.close()
        write_mini_pack(self.root / "火影-003", "规格C")
        again = sync_workbook(self.root)
        rows = {row["商品标识*"]: row for row in self._titles(again["path"])}
        self.assertEqual(rows["火影-001"]["商品标题*"], "卡游手改标题不要覆盖")
        self.assertIn("火影-003", rows)
        self.assertEqual(rows["火影-001"]["价格*"], 9.9)

    def test_missing_folder_marked_not_deleted(self):
        write_mini_pack(self.root / "火影-001")
        write_mini_pack(self.root / "火影-002")
        sync_workbook(self.root, {"品牌": "卡游"})
        shutil.rmtree(self.root / "火影-002")
        synced = sync_workbook(self.root)
        rows = {row["商品标识*"]: row for row in self._titles(synced["path"])}
        self.assertIn("火影-002", rows)
        self.assertEqual(rows["火影-002"]["处理状态"], MISSING_STATUS)
        self.assertIn("火影-001", rows)

    def test_folder_without_images_is_error(self):
        emptyish = self.root / "空图夹"
        emptyish.mkdir()
        (emptyish / "readme.txt").write_text("no images", encoding="utf-8")
        scan = scan_workspace(self.root)
        self.assertTrue(any(item["folder"] == "空图夹" for item in scan["errors"]))
        self.assertEqual(scan["folders"], [])

    def test_ignores_xlsx_inside_pack(self):
        pack = write_mini_pack(self.root / "火影-001")
        (pack / "夹内清单.xlsx").write_bytes(b"not-a-real-xlsx")
        scan = scan_workspace(self.root)
        self.assertEqual(len(scan["folders"]), 1)
        self.assertEqual(scan["folders"][0]["product_id"], "火影-001")
        self.assertEqual(scan["errors"], [])

    def test_skips_hidden_and_empty(self):
        (self.root / ".hidden").mkdir()
        write_mini_pack(self.root / ".hidden")
        (self.root / "空白").mkdir()
        scan = scan_workspace(self.root)
        reasons = {item["folder"]: item["reason"] for item in scan["skipped"]}
        self.assertEqual(reasons.get(".hidden"), "隐藏目录")
        self.assertEqual(reasons.get("空白"), "空目录")

    def test_long_folder_name_clipped_to_title(self):
        name = "测" * 31
        write_mini_pack(self.root / name)
        synced = sync_workbook(self.root, {"品牌": ""})
        row = self._titles(synced["path"])[0]
        self.assertLessEqual(workspace.title_width(row["商品标题*"]), 60)
        self.assertEqual(row["商品标题*"], clip_title(name))
        self.assertEqual(row["型号"], name)

    def test_sku_overrides_kept_on_rescan(self):
        write_mini_pack(self.root / "火影-001", "原名")
        synced = sync_workbook(self.root, {"品牌": "卡游"})
        path = Path(synced["path"])
        wb = load_workbook(path)
        sku = wb["SKU规格"]
        sku.cell(2, 3).value = "手改规格"
        sku.cell(2, 4).value = 8.8
        wb.save(path)
        wb.close()
        again = sync_workbook(self.root)
        wb = load_workbook(again["path"], data_only=True)
        try:
            self.assertEqual(wb["SKU规格"].cell(2, 3).value, "手改规格")
            self.assertEqual(wb["SKU规格"].cell(2, 4).value, 8.8)
        finally:
            wb.close()


class LiveWorkspacePacksTests(unittest.TestCase):
    def test_two_real_packs_validate(self):
        packs = list_naruto_packs()
        if len(packs) < 2:
            self.skipTest("需要 testdata 里两个火影真图夹")
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        try:
            link_folder(packs[0], root / "火影-001")
            link_folder(packs[1], root / "火影-002")
        except OSError as exc:
            self.skipTest(f"无法连接真图夹: {exc}")
        synced = sync_workbook(root, {
            "品牌": "卡游",
            "商品属性模板": "中性笔",
            "物流模板": "48小时",
            "销售模板": "仓库多规格",
            "价格": 9.9,
            "库存": 20,
        })
        import importlib.util
        spec = importlib.util.spec_from_file_location("seller", Path(__file__).with_name("千牛自动上架.py"))
        seller = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(seller)
        results = seller.validate_workbook(synced["path"])
        self.assertGreaterEqual(len(results), 2)
        self.assertTrue(all(item["validation"] == "通过" for item in results[:2]), results[0].get("errors"))
        self.assertTrue(results[0]["product"].get("pack_dir"))
        self.assertNotEqual(results[0]["product"]["title"], results[1]["product"]["title"])


if __name__ == "__main__":
    unittest.main()
