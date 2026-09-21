import argparse, json, os, re, sys, uuid, importlib, importlib.util
from io import BytesIO
from collections import Counter
from pathlib import Path
from datetime import datetime
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REQUIRED = ["商品标识*", "商品标题*", "类目*", "主图路径*", "价格*", "库存*"]
NEW_REQUIRED = [
    "商品标识*", "商品标题*", "价格*", "库存*",
    "图片包路径*", "商品属性模板*", "物流模板*", "销售模板*",
]
TEMPLATE_MARKERS = {"商品属性模板*", "物流模板*", "销售模板*", "图片包路径*"}
NEUTRAL_PEN_REQUIRED_ATTRIBUTES = {
    "笔头类型", "笔芯颜色", "闭合方式", "风格", "功能", "品牌", "型号",
    "适用场景", "适用人群", "包装方式", "采购地",
}
CORRECTION_TAPE_REQUIRED_ATTRIBUTES = {"附加功能", "风格", "可换芯设计", "适用场景"}
CATEGORY_ATTRIBUTES = {
    "中性笔": NEUTRAL_PEN_REQUIRED_ATTRIBUTES,
    "修正带": CORRECTION_TAPE_REQUIRED_ATTRIBUTES,
}
SHIPPING_TIMES = {"24小时内发货提升转化", "48小时内发货", "大于48小时发货"}
STOCK_DEDUCTION = {"拍下减库存", "付款减库存"}
TEMPLATE_HEADERS = [
    "商品标识*", "商品标题*", "品牌*", "型号", "图片包路径*",
    "商品属性模板*", "物流模板*", "销售模板*",
    "价格*", "库存*", "运费模板", "导购标题", "货号",
    "处理状态", "错误信息",
]
SKU_HEADERS = ["商品标识", "色号", "规格名称", "价格", "库存"]
NARUTO_PACK_NAME = "卡游火影忍者中性笔盲盒忍道版第1弹隐藏款鸣人水门动漫文具周边"
ANSI_ESCAPE = re.compile(r"\x1b\[[0-9;]*[A-Za-z]|\x1b\][^\x07]*\x07|\x1b[@-Z\\-_]")


def excel_safe(value):
    if value is None or isinstance(value, (int, float)):
        return value
    text = ANSI_ESCAPE.sub("", str(value))
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > 32000:
        text = text[:32000]
    return text


def _load_parser():
    spec = importlib.util.spec_from_file_location("qianniu_parse", Path(__file__).with_name("商品解析.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


_parser = _load_parser()
scan_image_pack = _parser.scan_image_pack
load_templates = _parser.load_templates
load_registry = _parser.load_registry
resolve_product = _parser.resolve_product
TEMPLATES_DIR = _parser.TEMPLATES_DIR


def present(value):
    return value is not None and str(value).strip() != ""


def valid_number(value, stock=False):
    if isinstance(value, bool):
        return False
    try:
        n = Decimal(str(value))
        return n.is_finite() and (n >= 0 and n == n.to_integral_value() if stock else n > 0)
    except (InvalidOperation, TypeError, ValueError, OverflowError):
        return False


def parse_json(value, field, row):
    try:
        return json.loads(value) if value not in (None, "") else None
    except Exception as exc:
        raise ValueError(f"第{row}行字段{field}不是有效JSON: {exc}")


def title_width(text):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text or ""))


def is_example_product(data):
    marker = str(data.get("商品标识*", "")).strip() + str(data.get("商品标题*", ""))
    return "示例" in marker


def uses_templates(data):
    return any(present(data.get(field)) for field in TEMPLATE_MARKERS)


def required_headers(headers):
    if TEMPLATE_MARKERS & set(headers):
        needed = list(NEW_REQUIRED)
        if "品牌*" not in headers and "品牌" not in headers:
            needed.append("品牌*")
        return needed
    return list(REQUIRED)


def naruto_pack_dir():
    return Path(__file__).resolve().parent / "testdata" / NARUTO_PACK_NAME


def relative_pack_path(workbook_path):
    pack = naruto_pack_dir()
    try:
        return Path(os.path.relpath(pack, Path(workbook_path).resolve().parent)).as_posix()
    except ValueError:
        return str(pack)


def _json_skus(skus):
    dump = []
    for sku in skus or []:
        item = {"规格": sku.get("规格") or {}, "价格": sku.get("价格", sku.get("price")), "库存": sku.get("库存", sku.get("stock"))}
        if present(sku.get("商家编码")):
            item["商家编码"] = sku.get("商家编码")
        dump.append(item)
    return json.dumps(dump, ensure_ascii=False)


def hydrate_template_row(data, sku_rows, base_dir, errors):
    names = {}
    if present(data.get("商品属性模板*")):
        names["attributes"] = str(data["商品属性模板*"]).strip()
    if present(data.get("物流模板*")):
        names["logistics"] = str(data["物流模板*"]).strip()
    if present(data.get("销售模板*")):
        names["sales"] = str(data["销售模板*"]).strip()
    templates = load_templates(names) if names else {}
    pack = None
    if present(data.get("图片包路径*")):
        folder = resolve_path(data["图片包路径*"], base_dir)
        if folder.is_dir():
            pack = scan_image_pack(folder)
        else:
            errors.append(f"图片包路径*: 目录不存在: {folder}")
    product = resolve_product(data, sku_rows, pack, templates, base_dir)
    if not present(data.get("类目*")) and product.get("category"):
        data["类目*"] = product["category"]
    if not present(data.get("品牌")) and product.get("brand"):
        data["品牌"] = product["brand"]
    if not present(data.get("商品属性(JSON)")) and product.get("attributes"):
        data["商品属性(JSON)"] = json.dumps(product["attributes"], ensure_ascii=False)
    if not present(data.get("主图路径*")) and product.get("main_images"):
        data["主图路径*"] = "|".join(str(path) for path in product["main_images"])
    if not present(data.get("3:4主图路径")) and product.get("portrait_images"):
        data["3:4主图路径"] = "|".join(str(path) for path in product["portrait_images"])
    if not present(data.get("详情图路径")) and product.get("detail_images"):
        data["详情图路径"] = "|".join(str(path) for path in product["detail_images"])
    if not present(data.get("发货时间")) and product.get("ship_time"):
        data["发货时间"] = product["ship_time"]
    if not present(data.get("发货地")) and product.get("ship_from"):
        data["发货地"] = product["ship_from"]
    if not present(data.get("运费模板")) and product.get("freight"):
        data["运费模板"] = product["freight"]
    if not present(data.get("库存扣减方式")) and product.get("stock_deduction"):
        data["库存扣减方式"] = product["stock_deduction"]
    if not present(data.get("上架时间")) and product.get("listing_time"):
        data["上架时间"] = product["listing_time"]
    dimensions = product.get("dimensions") or {}
    if dimensions and not present(data.get("规格1名称")):
        name, options = next(iter(dimensions.items()))
        data["规格1名称"] = name
        data["规格1值(JSON)"] = json.dumps(options, ensure_ascii=False)
    if product.get("skus") and not present(data.get("SKU明细(JSON)")) and not present(data.get("SKU明细(JSON)*")):
        data["SKU明细(JSON)"] = _json_skus(product["skus"])
    return product


def category_profile(category):
    text = str(category or "")
    for name, required in CATEGORY_ATTRIBUTES.items():
        if name in text:
            return name, required
    return "", set()


def split_paths(value):
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def resolve_path(raw, base_dir):
    path = Path(raw.strip())
    return path if path.is_absolute() else Path(base_dir or Path.cwd()) / path


def inspect_images(value, field, base_dir, kind, errors):
    paths = split_paths(value)
    if kind != "detail" and len(paths) > 5:
        errors.append(f"{field}: 最多5张，当前{len(paths)}张")
    for raw in paths:
        path = resolve_path(raw, base_dir)
        if not path.is_file():
            errors.append(f"{field}: 文件不存在或不是文件: {path}")
            continue
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            errors.append(f"{field}: 非工具支持的扩展名: {path.suffix}")
            continue
        try:
            from PIL import Image
            with Image.open(path) as picture:
                picture.load()
                width, height = picture.size
        except Exception as exc:
            errors.append(f"{field}: 图片无法解码: {path}: {exc}")
            continue
        if kind == "square" and width != height:
            errors.append(f"{field}: 1:1主图必须是正方形: {path.name} {width}x{height}")
        elif kind == "portrait" and not (0.7 <= (width / height if height else 0) <= 0.8):
            errors.append(f"{field}: 3:4主图比例应为3:4: {path.name} {width}x{height}")
        elif kind == "detail" and height / max(width, 1) > 2:
            errors.append(f"{field}: 详情图高宽比不能大于2: {path.name} {width}x{height}")
    return [resolve_path(raw, base_dir) for raw in paths if resolve_path(raw, base_dir).is_file()]


def validate_row(headers, values, row, base_dir=None, sku_rows=None):
    data = dict(zip(headers, values))
    errors = []
    if uses_templates(data):
        try:
            hydrate_template_row(data, sku_rows or [], base_dir, errors)
        except (ValueError, FileNotFoundError, json.JSONDecodeError) as exc:
            errors.append(str(exc))
    if not present(data.get("品牌")) and present(data.get("品牌*")):
        data["品牌"] = data["品牌*"]
    needed = NEW_REQUIRED if uses_templates(data) or (TEMPLATE_MARKERS & set(headers)) else REQUIRED
    if TEMPLATE_MARKERS & set(headers) and not present(data.get("品牌*") or data.get("品牌")):
        if "品牌*" not in needed:
            errors.append("缺少品牌*")
    for field in needed:
        if field == "主图路径*" and present(data.get("图片包路径*")):
            continue
        if field == "类目*" and present(data.get("商品属性模板*")):
            continue
        if not present(data.get(field)):
            errors.append(f"缺少{field}")
    if uses_templates(data) and present(data.get("图片包路径*")) and not present(data.get("主图路径*")):
        if not any(str(item).startswith("图片包路径*") for item in errors):
            errors.append("主图路径*: 图片包中没有可用的1:1主图")
    parsed_fields = {}
    if present(data.get("SKU明细(JSON)")) and present(data.get("SKU明细(JSON)*")):
        errors.append("SKU明细(JSON): 新旧SKU列不能同时填写，请仅保留一列数据")
    for field in ("商品属性(JSON)", "规格1值(JSON)", "规格2值(JSON)", "规格3值(JSON)", "SKU明细(JSON)*", "SKU明细(JSON)"):
        if data.get(field) not in (None, ""):
            try:
                parsed = parse_json(data[field], field, row)
                expected = dict if field == "商品属性(JSON)" else list
                if not isinstance(parsed, expected):
                    errors.append(f"{field}: 必须是{'对象' if expected is dict else '数组'}")
                else:
                    parsed_fields[field] = parsed
            except ValueError as exc:
                errors.append(str(exc))
    name, required_attrs = category_profile(data.get("类目*", ""))
    attrs = parsed_fields.get("商品属性(JSON)")
    if required_attrs:
        if not isinstance(attrs, dict):
            errors.append(f"商品属性(JSON): {name}类目必须提供属性对象")
        else:
            missing = sorted(required_attrs - set(attrs))
            if missing:
                errors.append(f"商品属性(JSON): 缺少{name}必填属性: " + "、".join(missing))
        for field in ("发货时间", "发货地", "运费模板"):
            if not present(data.get(field)):
                errors.append(f"{field}: {name}类目页面核验为必填")
    if present(data.get("发货时间")) and str(data["发货时间"]).strip() not in SHIPPING_TIMES:
        errors.append("发货时间: 只允许 24小时内发货提升转化 / 48小时内发货 / 大于48小时发货")
    if present(data.get("库存扣减方式")) and str(data["库存扣减方式"]).strip() not in STOCK_DEDUCTION:
        errors.append("库存扣减方式: 只允许 拍下减库存 或 付款减库存")
    title = str(data.get("商品标题*", "") or "")
    if present(title) and title_width(title) > 60:
        errors.append("商品标题*: 超过30个汉字（60字符）")
    brand = ""
    if isinstance(attrs, dict) and present(attrs.get("品牌")):
        brand = str(attrs.get("品牌")).strip()
    elif present(data.get("品牌")):
        brand = str(data.get("品牌")).strip()
    if brand and present(data.get("品牌")) and str(data.get("品牌")).strip() != brand:
        errors.append("品牌: 与商品属性中的品牌不一致")
    if brand and brand not in title:
        errors.append("商品标题*: 必须包含品牌属性，避免标题与品牌不一致")
    for field in ("价格*", "库存*"):
        if present(data.get(field)) and not valid_number(data[field], field == "库存*"):
            errors.append(f"{field}: 必须为非负整数" if field == "库存*" else f"{field}: 必须为大于零的有限数字")
    dimensions = {}
    for i in (1, 2, 3):
        spec_name, options = data.get(f"规格{i}名称"), parsed_fields.get(f"规格{i}值(JSON)", [])
        if present(spec_name) or options:
            if not present(spec_name) or not options or any(not isinstance(v, str) or not v.strip() for v in options):
                errors.append(f"规格{i}: 名称和值数组必须同时提供，值必须为非空字符串")
            elif str(spec_name) in dimensions or len(options) != len(set(options)):
                errors.append(f"规格{i}: 名称或选项重复")
            else:
                dimensions[str(spec_name)] = options
    skus = parsed_fields.get("SKU明细(JSON)", parsed_fields.get("SKU明细(JSON)*", []))
    if dimensions and not skus:
        errors.append("SKU明细(JSON): 多规格商品必须提供SKU")
    seen = set()
    for i, sku in enumerate(skus, 1):
        field = f"SKU明细(JSON)[{i}]"
        if not isinstance(sku, dict):
            errors.append(f"{field}: 必须是对象")
            continue
        specs = sku.get("规格", {})
        if not isinstance(specs, dict) or any(not isinstance(v, str) or not v.strip() for v in specs.values()):
            errors.append(f"{field}.规格: 必须为字符串值对象")
        else:
            key = json.dumps(specs, sort_keys=True, ensure_ascii=False)
            if key in seen:
                errors.append(f"{field}.规格: 组合重复")
            seen.add(key)
            if set(specs) != set(dimensions) or any(v not in dimensions.get(k, []) for k, v in specs.items()):
                errors.append(f"{field}.规格: 必须与规格名称和选项一致")
        for label in ("价格", "库存"):
            if not valid_number(sku.get(label), label == "库存"):
                errors.append(f"{field}.{label}: 数值无效（价格须为有限正数，库存须为非负整数）")
    inspect_images(data.get("主图路径*"), "主图路径*", base_dir, "square", errors)
    inspect_images(data.get("3:4主图路径"), "3:4主图路径", base_dir, "portrait", errors)
    inspect_images(data.get("详情图路径"), "详情图路径", base_dir, "detail", errors)
    if present(data.get("上架时间")) and data["上架时间"] != "放入仓库":
        errors.append("上架时间: 只允许放入仓库")
    if is_example_product(data):
        errors.append("示例数据禁止提交真实店铺，请替换为真实商品资料")
    return data, errors


def validate_product(product, base_dir=None):
    errors = []
    title = str(product.get("title") or "")
    attrs = product.get("attributes") or {}
    brand = str(product.get("brand") or "").strip()
    if isinstance(attrs, dict) and present(attrs.get("品牌")):
        brand = str(attrs.get("品牌")).strip()
    name, required_attrs = category_profile(product.get("category") or "")
    if required_attrs:
        if not isinstance(attrs, dict):
            errors.append(f"商品属性(JSON): {name}类目必须提供属性对象")
        else:
            missing = sorted(required_attrs - set(attrs))
            if missing:
                errors.append(f"商品属性(JSON): 缺少{name}必填属性: " + "、".join(missing))
    if not brand:
        errors.append("品牌: 必须填写")
    elif brand not in title:
        errors.append("商品标题*: 必须包含品牌属性，避免标题与品牌不一致")
    if present(title) and title_width(title) > 60:
        errors.append("商品标题*: 超过30个汉字（60字符）")
    listing = str(product.get("listing_time") or product.get("上架时间") or "").strip()
    if present(listing) and listing != "放入仓库":
        errors.append("上架时间: 只允许放入仓库")
    mains = product.get("main_images") or []
    if not mains:
        errors.append("主图路径*: 至少需要1张1:1主图")
    inspect_images("|".join(str(path) for path in mains), "主图路径*", base_dir, "square", errors)
    inspect_images("|".join(str(path) for path in (product.get("portrait_images") or [])), "3:4主图路径", base_dir, "portrait", errors)
    inspect_images("|".join(str(path) for path in (product.get("detail_images") or [])), "详情图路径", base_dir, "detail", errors)
    return errors


def row_to_product(data, base_dir, sku_rows=None):
    if uses_templates(data) or present(data.get("图片包路径*")):
        names = {}
        if present(data.get("商品属性模板*")):
            names["attributes"] = str(data["商品属性模板*"]).strip()
        if present(data.get("物流模板*")):
            names["logistics"] = str(data["物流模板*"]).strip()
        if present(data.get("销售模板*")):
            names["sales"] = str(data["销售模板*"]).strip()
        templates = load_templates(names) if names else {}
        pack = None
        if present(data.get("图片包路径*")):
            folder = resolve_path(data["图片包路径*"], base_dir)
            if folder.is_dir():
                pack = scan_image_pack(folder)
        return resolve_product(data, sku_rows, pack, templates, base_dir)
    parsed = {}
    for field in ("商品属性(JSON)", "规格1值(JSON)", "规格2值(JSON)", "规格3值(JSON)", "SKU明细(JSON)*", "SKU明细(JSON)"):
        if data.get(field) not in (None, ""):
            parsed[field] = json.loads(data[field]) if isinstance(data[field], str) else data[field]
    dimensions = {}
    for i in (1, 2, 3):
        name = data.get(f"规格{i}名称")
        options = parsed.get(f"规格{i}值(JSON)") or []
        if present(name) and options:
            dimensions[str(name)] = options
    brand = str(data.get("品牌") or data.get("品牌*") or "").strip()
    attrs = parsed.get("商品属性(JSON)") or {}
    if isinstance(attrs, dict) and present(attrs.get("品牌")):
        brand = str(attrs.get("品牌")).strip()
    return {
        "product_id": str(data.get("商品标识*", "")).strip(),
        "title": str(data.get("商品标题*", "")).strip(),
        "guide_title": str(data.get("导购标题") or "").strip(),
        "category": str(data.get("类目*", "")).strip(),
        "brand": brand,
        "outer_id": str(data.get("货号") or data.get("商家编码") or "").strip(),
        "attributes": attrs if isinstance(attrs, dict) else {},
        "dimensions": dimensions,
        "skus": parsed.get("SKU明细(JSON)", parsed.get("SKU明细(JSON)*", [])) or [],
        "price": data.get("价格*"),
        "stock": data.get("库存*"),
        "main_images": [resolve_path(p, base_dir) for p in split_paths(data.get("主图路径*"))],
        "portrait_images": [resolve_path(p, base_dir) for p in split_paths(data.get("3:4主图路径"))],
        "detail_images": [resolve_path(p, base_dir) for p in split_paths(data.get("详情图路径"))],
        "description": str(data.get("商品描述") or "").strip(),
        "freight": str(data.get("运费模板") or "").strip(),
        "ship_time": str(data.get("发货时间") or "").strip(),
        "ship_from": str(data.get("发货地") or "").strip(),
        "stock_deduction": str(data.get("库存扣减方式") or "").strip(),
    }


def read_sku_rows(wb):
    if "SKU规格" not in wb.sheetnames:
        return []
    ws = wb["SKU规格"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(x or "").strip() for x in next(rows, ())]
    items = []
    for values in rows:
        data = dict(zip(headers, values))
        if any(present(data.get(field)) for field in SKU_HEADERS):
            items.append(data)
    return items


def read_workbook(path):
    wb = load_workbook(BytesIO(Path(path).read_bytes()), data_only=True, read_only=True)
    try:
        ws = wb["商品清单"] if "商品清单" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = [str(x or "").strip() for x in next(rows, ())]
        needed = required_headers(headers)
        missing = [h for h in needed if h not in headers]
        duplicates = [h for h, n in Counter(headers).items() if h and n > 1]
        if missing or duplicates:
            raise ValueError(f"模板缺列: {missing}; 重复列: {duplicates}")
        meta = {"处理状态", "错误信息", "上架时间", "库存扣减方式"}
        products = []
        for row_num, values in enumerate(rows, start=2):
            if any(present(v) for h, v in zip(headers, values) if h and h not in meta):
                products.append((row_num, headers, values))
        return products, read_sku_rows(wb)
    finally:
        wb.close()


def read_products(path):
    rows, _ = read_workbook(path)
    yield from rows


def write_log(path, entries, output=None):
    source = Path(path).resolve()
    target = Path(output).resolve() if output else source.with_name(source.stem + ".校验结果.xlsx")
    if target == source:
        raise ValueError("结果文件不能覆盖输入文件")
    wb = load_workbook(path)
    temp = target.with_name(target.stem + "." + uuid.uuid4().hex + ".tmp.xlsx")
    try:
        ws = wb["处理日志"] if "处理日志" in wb.sheetnames else wb.create_sheet("处理日志")
        previous = [list(row) for row in ws.iter_rows(min_row=2, values_only=True) if any(present(v) for v in row)]
        ws.delete_rows(1, ws.max_row)
        ws.append(["时间", "Excel行号", "商品标识", "结果", "失败字段", "页面提示", "详情", "淘宝商品ID"])
        for entry in previous + entries:
            ws.append([excel_safe(v) for v in entry])
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(temp)
        os.replace(temp, target)
    finally:
        wb.close()
        if temp.exists():
            temp.unlink()
    return target


def validate_workbook(path):
    rows, sku_rows = read_workbook(path)
    ids = Counter(str(dict(zip(h, v)).get("商品标识*", "")).strip() for _, h, v in rows)
    results = []
    base_dir = Path(path).resolve().parent
    for row, headers, values in rows:
        data = dict(zip(headers, values))
        product_id = str(data.get("商品标识*", "")).strip()
        related = [item for item in sku_rows if not present(item.get("商品标识")) or str(item.get("商品标识")).strip() == product_id]
        data, errors = validate_row(headers, values, row, base_dir, related)
        product_id = str(data.get("商品标识*", "")).strip()
        if ids[product_id] > 1:
            errors.append("商品标识*: 整批重复，所有重复行均不可执行")
        item = {
            "row": row,
            "product_id": product_id,
            "validation": "失败" if errors else "通过",
            "execution": "未执行",
            "notice": "",
            "errors": errors,
            "product": row_to_product(data, base_dir, related),
        }
        results.append(item)
    try:
        import job_session
        job_session.merge_execution(results, path)
        job_session.remember_workbook(path, results)
    except Exception:
        pass
    return results


def load_executor():
    existing = sys.modules.get("千牛网页执行")
    if existing is not None:
        return existing
    sidecar = Path(__file__).with_name("千牛网页执行.py")
    if sidecar.is_file():
        spec = importlib.util.spec_from_file_location("千牛网页执行", sidecar)
        module = importlib.util.module_from_spec(spec)
        sys.modules["千牛网页执行"] = module
        spec.loader.exec_module(module)
        return module
    return importlib.import_module("千牛网页执行")


def load_web_fill():
    root = Path(__file__).resolve().parent
    pipeline_file = root / "web_fill" / "pipeline.py"
    if not pipeline_file.is_file():
        pipeline_file = Path(sys.executable).resolve().parent / "web_fill" / "pipeline.py"
    if pipeline_file.is_file():
        parent = str(pipeline_file.parent.parent)
        if parent not in sys.path:
            sys.path.insert(0, parent)
        for key in list(sys.modules):
            if key == "web_fill" or key.startswith("web_fill."):
                sys.modules.pop(key, None)
        import web_fill
        return web_fill
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    import web_fill
    return web_fill


def run_web_batch(products, confirm_submit=False, limit=None):
    return load_web_fill().run_batch(products, confirm_submit=confirm_submit, limit=limit)


def _style_header(sheet, fill, font, wrap, thin):
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment, cell.border = fill, font, wrap, thin
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}1"


def _write_option_column(sheet, column, values):
    for index, value in enumerate(values, 1):
        sheet.cell(index, column, value)
    last = max(len(values), 1)
    return f"=_模板选项!${get_column_letter(column)}$1:${get_column_letter(column)}${last}"


FIELD_NOTES = [
    ["商品标识", "工具必填", "火影-001", "日志关联用；不要写「示例」"],
    ["商品标题", "平台必填", "最多30汉字", "必须包含品牌全文，文件夹全名超长时请缩短"],
    ["品牌 / 型号", "品牌必填", "卡游 / 忍道版第1弹", "品牌、型号、标题不进分类模板，每条商品自填"],
    ["图片包路径", "必填", "扁平文件夹", "不要再手填主图路径。文件名规则如下"],
    ["图片包·1:1主图", "平台必填", "宝贝主图01.jpg …", "按文件名排序，最多用5张"],
    ["图片包·3:4主图", "非必填", "文件名含 3比4 / 3-4", "缺则跳过；也可按接近 3:4 的像素归入"],
    ["图片包·详情图", "建议", "详情01.jpg …", "按文件名排序；扩展名 jpg/png/webp 均可"],
    ["图片包·SKU图", "多规格必填", "颜色01-规格名.jpg", "横杠后的文字就是默认规格名，可在 SKU规格 表改"],
    ["商品属性模板", "必填", "下拉：中性笔", "类目和必填属性默认值在 templates/attributes/，不用写属性 JSON"],
    ["物流模板", "必填", "下拉：48小时 / 24小时", "发货时间、发货地、运费模板名（文具用品 包邮）"],
    ["销售模板", "必填", "下拉：仓库多规格", "单层商品规格；书写粗细默认 0.05mm；放入仓库"],
    ["价格 / 库存", "平台必填", "正数 / 非负整数", "作为各 SKU 默认值；SKU规格 表可按行覆盖"],
    ["运费模板 / 导购标题 / 货号", "可选覆盖", "留空用模板", "同行同名字段非空则覆盖物流/销售模板"],
    ["SKU规格表", "可选", "商品标识+色号", "色号如 颜色01；规格名称为空则用文件名；整表缺行则 16 色都用文件名"],
    ["旧列兼容", "读旧表", "商品属性(JSON) / SKU明细(JSON) / 主图路径*", "旧清单仍可校验，不必改成图片包"],
    ["上架时间", "固定", "放入仓库", "不会执行立刻上架或定时上架"],
]


def _listing_cell(row, field):
    if field not in (row or {}):
        return ""
    value = row.get(field)
    return "" if value is None else value


def write_listing_workbook(path, product_rows=None, sku_rows=None, log_rows=None):
    """写出带下拉的商品清单 / SKU规格 / 字段说明 / 处理日志。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    registry = load_registry()
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", name="Microsoft YaHei", size=10)
    wrap = Alignment(wrap_text=True, vertical="center")
    thin = Border(*(Side(style="thin", color="D9E2F3") for _ in range(4)))
    options = wb.active
    options.title = "_模板选项"
    attr_formula = _write_option_column(options, 1, registry.get("attributes") or ["中性笔"])
    log_formula = _write_option_column(options, 2, registry.get("logistics") or ["48小时", "24小时"])
    sales_formula = _write_option_column(options, 3, registry.get("sales") or ["仓库多规格"])
    options.sheet_state = "hidden"
    sheet = wb.create_sheet("商品清单", 0)
    sheet.append(TEMPLATE_HEADERS)
    for item in product_rows or []:
        sheet.append([_listing_cell(item, field) for field in TEMPLATE_HEADERS])
    _style_header(sheet, header_fill, header_font, wrap, thin)
    for index, width in enumerate([12, 34, 10, 16, 42, 16, 12, 14, 10, 8, 16, 16, 12, 10, 20], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    columns = {name: index for index, name in enumerate(TEMPLATE_HEADERS, 1)}
    for formula, field in (
        (attr_formula, "商品属性模板*"),
        (log_formula, "物流模板*"),
        (sales_formula, "销售模板*"),
    ):
        letter = get_column_letter(columns[field])
        rule = DataValidation(type="list", formula1=formula, allow_blank=False, showErrorMessage=True)
        rule.error = "请从下拉列表选择模板"
        rule.errorTitle = "模板不在清单中"
        rule.add(f"{letter}2:{letter}1000")
        sheet.add_data_validation(rule)
    sku_sheet = wb.create_sheet("SKU规格")
    sku_sheet.append(SKU_HEADERS)
    for item in sku_rows or []:
        sku_sheet.append([_listing_cell(item, field) for field in SKU_HEADERS])
    _style_header(sku_sheet, header_fill, header_font, wrap, thin)
    for index, width in enumerate([12, 12, 36, 10, 8], 1):
        sku_sheet.column_dimensions[get_column_letter(index)].width = width
    notes = wb.create_sheet("字段说明")
    notes.append(["字段", "是否必填", "格式/示例", "说明"])
    for row in FIELD_NOTES:
        notes.append(row)
    for index, width in enumerate([22, 12, 36, 56], 1):
        notes.column_dimensions[get_column_letter(index)].width = width
    _style_header(notes, header_fill, header_font, wrap, thin)
    log = wb.create_sheet("处理日志")
    log.append(["时间", "Excel行号", "商品标识", "结果", "失败字段", "页面提示", "详情", "淘宝商品ID"])
    for entry in log_rows or []:
        log.append([excel_safe(value) for value in entry])
    _style_header(log, header_fill, header_font, wrap, thin)
    wb.save(path)
    wb.close()
    return path


def create_template(path):
    pack_rel = relative_pack_path(path)
    pack = naruto_pack_dir()
    scanned = scan_image_pack(pack) if pack.is_dir() else {"skus": []}
    product_rows = [{
        "商品标识*": "火影-001",
        "商品标题*": "卡游火影忍者中性笔盲盒忍道版",
        "品牌*": "卡游",
        "型号": "忍道版第1弹",
        "图片包路径*": pack_rel,
        "商品属性模板*": "中性笔",
        "物流模板*": "48小时",
        "销售模板*": "仓库多规格",
        "价格*": 9.9,
        "库存*": 20,
        "运费模板": "",
        "导购标题": "",
        "货号": "",
        "处理状态": "待处理",
        "错误信息": "",
    }]
    sku_rows = [
        {"商品标识": "火影-001", "色号": item.get("slot"), "规格名称": "", "价格": "", "库存": ""}
        for item in scanned.get("skus") or []
    ]
    return write_listing_workbook(path, product_rows, sku_rows)


def main():
    ap = argparse.ArgumentParser(description="千牛商品批量录入（默认只校验，不提交）")
    ap.add_argument("excel", nargs="?", help="千牛商品清单模板.xlsx")
    ap.add_argument("--submit", action="store_true", help="连接已登录千牛窗口并填写，默认仍不点击提交")
    ap.add_argument("--confirm-submit", action="store_true", help="在已选放入仓库后点击提交宝贝信息")
    ap.add_argument("--limit", type=int, help="本次最多处理的通过校验商品数")
    ap.add_argument("--output", help="结果Excel路径，不能覆盖输入")
    ap.add_argument("--log", help="独立JSON报告路径，必须是新文件")
    ap.add_argument("--init-template", help="生成/覆盖商品清单模板")
    ap.add_argument("--workspace", help="工作空间目录：扫描一级子夹并生成/合并 商品清单.xlsx")
    args = ap.parse_args()
    if args.init_template:
        target = create_template(args.init_template)
        print(json.dumps({"template": str(target)}, ensure_ascii=False, indent=2))
        return 0
    if args.workspace:
        from workspace import sync_workbook
        synced = sync_workbook(args.workspace)
        source = Path(synced["path"]).resolve()
    elif args.excel:
        source = Path(args.excel).resolve()
    else:
        print("请提供Excel路径、--workspace，或使用 --init-template", file=sys.stderr)
        return 2
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    output = Path(args.output).resolve() if args.output else source.with_name(f"{source.stem}.结果-{stamp}.xlsx")
    journal = Path(args.log).resolve() if args.log else output.with_suffix(".json")
    if source == output or journal in (source, output):
        print("输入、Excel结果、JSON报告路径必须互不相同", file=sys.stderr)
        return 2
    try:
        results = validate_workbook(source)
    except Exception as exc:
        print(f"读取/校验失败: {exc}", file=sys.stderr)
        return 2
    executable = [r for r in results if r["validation"] == "通过"]
    if args.submit:
        if args.confirm_submit and not executable:
            print("没有可通过校验的真实商品，已拒绝提交", file=sys.stderr)
        elif executable:
            try:
                executed = run_web_batch(
                    [{**r["product"], "row": r["row"]} for r in executable],
                    confirm_submit=args.confirm_submit,
                    limit=args.limit,
                )
            except Exception as exc:
                print(f"网页执行失败: {exc}", file=sys.stderr)
                for item in results:
                    if item["validation"] == "通过":
                        item["execution"] = "失败"
                        item["notice"] = str(exc)
                        item["errors"] = item["errors"] + [str(exc)]
            else:
                by_row = {item.get("row"): item for item in executed}
                for item in results:
                    web = by_row.get(item["row"])
                    if not web:
                        continue
                    item["execution"] = web.get("execution", "失败")
                    item["notice"] = web.get("notice", "")
                    item["errors"] = item["errors"] + list(web.get("errors") or [])
                    for key in ("taobao_item_id", "view_url", "edit_url", "url"):
                        if web.get(key):
                            item[key] = web.get(key)
                    if item["errors"] and item["execution"] in {"未执行", "已填写未提交"}:
                        item["execution"] = web.get("execution") or "失败"
    valid = sum(r["validation"] == "通过" for r in results)
    report = {
        "source": str(source),
        "total": len(results),
        "valid": valid,
        "execution_enabled": bool(args.submit),
        "confirm_submit": bool(args.confirm_submit),
        "results": [{k: v for k, v in r.items() if k != "product"} for r in results],
    }
    try:
        journal.parent.mkdir(parents=True, exist_ok=True)
        with journal.open("x", encoding="utf-8") as stream:
            json.dump(report, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
    except OSError as exc:
        print(f"独立报告保存失败，停止操作: {exc}", file=sys.stderr)
        return 2
    entries = [[
        stamp, r["row"], r["product_id"],
        ("校验" + r["validation"]) if r["execution"] == "未执行" else r["execution"],
        "; ".join(e.split(":", 1)[0] for e in r["errors"]),
        r.get("notice") or "",
        "; ".join(r["errors"]) or (r.get("notice") or "仅本地校验通过；网页未执行"),
        r.get("taobao_item_id") or "",
    ] for r in results]
    try:
        write_log(source, entries, output)
    except Exception as exc:
        print(f"Excel结果保存失败: {exc}；完整结果已保留: {journal}", file=sys.stderr)
        return 2
    print(json.dumps({"total": len(results), "valid": valid, "report": str(journal), "output": str(output)}, ensure_ascii=False, indent=2))
    if not args.submit:
        print("当前为 dry-run：只校验，结果另存，未操作网页。")
        return 1 if valid != len(results) else 0
    if not executable:
        print("没有可执行的真实商品，未操作网页。")
        return 1
    if not args.confirm_submit:
        print("已尝试网页填写；未加 --confirm-submit，不会点击“提交宝贝信息”。")
    failed = [r for r in results if r["validation"] != "通过" or r["execution"] in {"失败", "暂停"}]
    return 1 if failed else 0

if __name__ == "__main__":
    sys.exit(main())
